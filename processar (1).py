"""
UNIASSELVI - Dashboard de Portfolios
Le as planilhas da pasta planilhas/ e gera saida/dashboard.html

PATCHES v2 aplicados:
  1. CH SEMANAL lida de 01_CONTROLE_TUTORIA.xlsx (não depende de LOTACAO)
  2. tem_lotacao = True somente quando há ch_semanal > 0 nos tutores
  3. status_ordem usa PERIODOS_ORDENS para datas de início corretas
  4. Situação do tutor corrigida quando nenhuma ordem está vencida ainda
  5. Campo 'sit' sincronizado após deduplicação por email
  6. Prints duplicados removidos de verificar_e_localizar()
  7. datas_por_tutor adicionado no formato antigo de agendas
"""

import pandas as pd
import json, os, sys, math, webbrowser, time, threading, glob
from pathlib import Path
from datetime import datetime, timezone, timedelta
from collections import defaultdict

PRAZOS_ORDENS = {
    'Ordem 1': '14/03/2026',
    'Ordem 2': '11/04/2026',
    'Ordem 3': '09/05/2026',
    'Ordem 4': '06/06/2026',
    'Ordem 5': '04/07/2026',
}
PERIODOS_ORDENS = {
    'Ordem 1': {'inicio': '16/02/2026', 'fim': '14/03/2026', 'semanas': 4},
    'Ordem 2': {'inicio': '16/03/2026', 'fim': '11/04/2026', 'semanas': 4},
    'Ordem 3': {'inicio': '13/04/2026', 'fim': '09/05/2026', 'semanas': 4},
    'Ordem 4': {'inicio': '11/05/2026', 'fim': '06/06/2026', 'semanas': 4},
    'Ordem 5': {'inicio': '08/06/2026', 'fim': '04/07/2026', 'semanas': 4},
}
CH_ADMIN_FATOR   = 0.25
CH_PRATICA_DURAC = 1.5


def _parse_ch(v):
    """PATCH 1 helper — converte CH SEMANAL (HH:MM ou decimal) para float horas."""
    if v is None: return None
    sv = str(v).strip()
    if sv in ('', 'nan', 'NaN', 'None', '0', '0.0'): return None
    try:
        if ':' in sv:
            parts = sv.split(':')
            result = float(parts[0]) + float(parts[1]) / 60
        else:
            result = float(sv.replace(',', '.'))
        return result if result > 0 else None
    except (ValueError, TypeError):
        return None


def achar_pasta_script():
    candidatos = []
    try:
        p = os.path.dirname(os.path.abspath(sys.argv[0]))
        if os.path.isdir(p): candidatos.append(p)
    except: pass
    try:
        p = os.path.dirname(os.path.abspath(__file__))
        if os.path.isdir(p): candidatos.append(p)
    except: pass
    try:
        p = os.getcwd()
        if os.path.isdir(p): candidatos.append(p)
    except: pass
    for p in candidatos:
        if os.path.isdir(os.path.join(p, "planilhas")): return p
        if os.path.isfile(os.path.join(p, "processar.py")): return p
    return candidatos[0] if candidatos else os.getcwd()


SCRIPT_DIR = achar_pasta_script()


def ler_url_file(path_url):
    try:
        with open(path_url, encoding='utf-8', errors='replace') as f:
            for line in f:
                if line.upper().startswith('URL='): return line[4:].strip()
    except: pass
    return None


def forcar_download_onedrive(path_url_file, destino, label):
    import subprocess, shutil, time
    path_xlsx = path_url_file.replace('.url', '').replace('.URL', '')
    try:
        subprocess.run(['attrib', '-P', '+U', path_url_file], capture_output=True, timeout=10)
    except Exception: pass
    for _ in range(6):
        if os.path.isfile(path_xlsx):
            with open(path_xlsx, 'rb') as f:
                header = f.read(4)
            if header == b'PK\x03\x04':
                shutil.copy2(path_xlsx, destino)
                print(f"  [OneDrive] Sincronizado: {label}")
                return destino
        time.sleep(5)
        print(f"  [OneDrive] Aguardando sync para {label}...")
    print(f"  [OneDrive] Timeout aguardando {label}.")
    return None


_KEYWORDS = {
    '01_CONTROLE_TUTORIA.xlsx': ['CONTROLE'],
    'PORTFOLIO_TUTOR.xlsx':     ['PORTFOLIO', 'PORTIFOLIO', 'PORTF'],
    'REL_GERAL_DE_GERENCIAMENTO.xlsx': ['GERENCIAMENTO', 'REL_GERAL'],
}
_ONEDRIVE_NAMES = {
    '01_CONTROLE_TUTORIA.xlsx': ['CONTROLE'],
    'PORTFOLIO_TUTOR.xlsx':     ['PORTF', 'PORTFOLIO'],
}


def _bate(caminho_arq, padrao):
    bn  = os.path.basename(caminho_arq).upper()
    kws = _KEYWORDS.get(padrao, [os.path.splitext(padrao)[0].upper()])
    return any(kw in bn for kw in kws)


def achar_arquivo(pasta, padrao):
    pasta_planilhas = os.path.join(pasta, "planilhas")
    direto = os.path.join(pasta_planilhas, padrao)
    if os.path.isfile(direto): return direto
    for arq in glob.glob(os.path.join(pasta_planilhas, "*.xls")) + glob.glob(os.path.join(pasta_planilhas, "*.xlsx")):
        if _bate(arq, padrao): return arq
    for arq in glob.glob(os.path.join(pasta_planilhas, "*.url")) + glob.glob(os.path.join(pasta_planilhas, "*.xlsx.url")) + glob.glob(os.path.join(pasta_planilhas, "*.xls.url")):
        if _bate(arq, padrao):
            url = ler_url_file(arq)
            if url:
                destino = os.path.join(pasta_planilhas, padrao)
                resultado = forcar_download_onedrive(arq, destino, padrao)
                if resultado: return resultado
    usuario = os.environ.get('USERNAME', os.environ.get('USER', 'leona'))
    for base in [
        f"C:\\Users\\{usuario}\\OneDrive - Uniasselvi",
        f"C:\\Users\\{usuario}\\OneDrive - UNIASSELVI",
        f"C:\\Users\\{usuario}\\OneDrive - Grupo Uniasselvi",
        f"C:\\Users\\{usuario}\\OneDrive",
    ]:
        if not os.path.isdir(base): continue
        for arq in glob.glob(os.path.join(base, "*.url")):
            if _bate(arq, padrao):
                url = ler_url_file(arq)
                if url:
                    destino = os.path.join(pasta_planilhas, padrao)
                    resultado = forcar_download_onedrive(arq, destino, padrao)
                    if resultado: return resultado
        for arq in glob.glob(os.path.join(base, "**", "*.xls"), recursive=True) + glob.glob(os.path.join(base, "**", "*.xlsx"), recursive=True):
            if _bate(arq, padrao): return arq
    return None


WATCH_MODE = len(sys.argv) > 1 and sys.argv[1].lower() == "watch"

CAT_MAP = {
    'ENF-INS (Multidisciplinar II)':
        'Multidisciplinar II - Enfermagem e Instrumentação Cirúrgica',
    'BIO-FISIO-EST-TO (Multidisciplinar III)':
        'Multidisciplinar III - Biomedicina Estética, Fisioterapia, Terapia Ocupacional e Estética e Cosmética',
    'BIO-BIO-FISIO-EST-TO (Multidisciplinar III)':
        'Multidisciplinar III - Biomedicina Estética, Fisioterapia, Terapia Ocupacional e Estética e Cosmética',
    'BIO-FAR (Multidisciplinar I)':
        'Multidisciplinar I - Biomedicina e Farmácia',
    'NUTRI (Multidisciplinar IV)':
        'Multidisciplinar IV - Nutrição',
    'QUÍMICA E FÍSICA':
        'Química e Física - Agronomia',
    'ENGMAKER':
        'EngeMaker | Química e Física - Engenharias e Licenciaturas',
    'ENGMAKER+QUÍMICA E FÍSICA':
        'EngeMaker | Química e Física - Engenharias e Licenciaturas',
}


def ts():
    BRT = timezone(timedelta(hours=-3))
    return datetime.now(BRT).strftime('%H:%M:%S')


def limpar(obj):
    if isinstance(obj, dict):   return {k: limpar(v) for k, v in obj.items()}
    if isinstance(obj, list):   return [limpar(v) for v in obj]
    if isinstance(obj, float) and math.isnan(obj): return None
    return obj


def verificar_e_localizar():
    pasta_planilhas = os.path.join(SCRIPT_DIR, "planilhas")
    os.makedirs(pasta_planilhas, exist_ok=True)
    print(f"  Script em : {SCRIPT_DIR}")
    print(f"  Planilhas : {pasta_planilhas}")
    print()
    cfg = {}
    cfg_file = os.path.join(SCRIPT_DIR, "config_links.json")
    if os.path.isfile(cfg_file):
        try:
            with open(cfg_file, encoding="utf-8") as f: cfg = json.load(f)
        except: pass
    cam_t = cfg.get("caminho_planilha_tutores", "").strip().strip('"')
    cam_p = cfg.get("caminho_planilha_portfolio", "").strip().strip('"')
    if cam_t and os.path.isfile(cam_t):
        p1 = cam_t; print(f"  [OK] {os.path.basename(p1)}")
    else:
        p1 = achar_arquivo(SCRIPT_DIR, "01_CONTROLE_TUTORIA.xlsx")
        if p1: print(f"  [OK] {os.path.basename(p1)}")
        else:  print(f"  [FALTA] 01_CONTROLE_TUTORIA.xlsx")
    if cam_p and os.path.isfile(cam_p):
        p2 = cam_p; print(f"  [OK] {os.path.basename(p2)}")
    else:
        p2 = achar_arquivo(SCRIPT_DIR, "PORTFOLIO_TUTOR.xlsx")
        if p2: print(f"  [OK] {os.path.basename(p2)}")
        else:  print(f"  [FALTA] PORTFOLIO_TUTOR.xlsx")
    # PATCH 6: prints duplicados de p1/p2 removidos aqui
    tmpl = os.path.join(SCRIPT_DIR, "template_dashboard.html")
    if os.path.isfile(tmpl): print(f"  [OK] template_dashboard.html")
    else:                    print(f"  [FALTA] template_dashboard.html")
    p3 = achar_arquivo(SCRIPT_DIR, "REL_GERAL_DE_GERENCIAMENTO.xlsx")
    if p3: print(f"  [OK] {os.path.basename(p3)}")
    else:  print(f"  [INFO] REL_GERAL_DE_GERENCIAMENTO.xlsx não encontrada (módulo desativado)")
    p4 = achar_arquivo(SCRIPT_DIR, "LOTACAO_TUTORES.xlsm") or achar_arquivo(SCRIPT_DIR, "LOTACAO_TUTORES.xlsx")
    if p4: print(f"  [OK] {os.path.basename(p4)}")
    else:  print(f"  [INFO] LOTACAO_TUTORES não encontrada (.xlsx/.xlsm)")
    # ── CSV de alunos por hub (igual aos outros arquivos: URL no Secret/env) ──
    p5 = None
    # 1. Tentar achar na pasta planilhas/ (já baixado anteriormente)
    p5 = achar_arquivo(SCRIPT_DIR, "Relatorio_alunos_por_hub.csv")
    if p5:
        print(f"  [OK] {os.path.basename(p5)}")
    else:
        # 2. Tentar baixar via variável de ambiente URL_ALUNOS_HUB (Secret GitHub)
        import re
        url_hub = os.environ.get("URL_ALUNOS_HUB", "").strip()
        if url_hub:
            print(f"  [Baixando] Relatorio_alunos_por_hub.csv via URL_ALUNOS_HUB...")
            try:
                import urllib.request
                # Converter link SharePoint/OneDrive para download direto
                # Tentar múltiplos formatos de URL
                def _build_dl_urls(url):
                    urls = []
                    if 'sharepoint.com' in url:
                        # Formato 1: adicionar &download=1
                        sep = '&' if '?' in url else '?'
                        urls.append(url + sep + 'download=1')
                        # Formato 2: download.aspx com token
                        m = re.search(r'/([A-Za-z0-9_-]{20,})[?]', url)
                        if m:
                            base = re.match(r'(https://[^/]+)', url).group(1)
                            user = re.search(r'/personal/([^/]+)/', url)
                            if user:
                                urls.append(f"{base}/personal/{user.group(1)}/_layouts/15/download.aspx?share={m.group(1)}")
                    elif '1drv.ms' in url:
                        sep = '&' if '?' in url else '?'
                        urls.append(url + sep + 'download=1')
                    urls.append(url)  # URL original como último recurso
                    return urls

                dest = os.path.join(pasta_planilhas, "Relatorio_alunos_por_hub.csv")
                downloaded = False
                for url_dl in _build_dl_urls(url_hub):
                    try:
                        req = urllib.request.Request(url_dl, headers={
                            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'})
                        with urllib.request.urlopen(req, timeout=120) as r:
                            data = r.read()
                        if len(data) > 10000 and b'<!DOCTYPE' not in data[:500]:
                            with open(dest, 'wb') as f_out: f_out.write(data)
                            p5 = dest
                            print(f"  [OK] Relatorio_alunos_por_hub.csv ({len(data):,} bytes)")
                            downloaded = True
                            break
                        else:
                            print(f"  [AVISO] URL retornou conteúdo inválido ({len(data)} bytes): {url_dl[:80]}")
                    except Exception as ex:
                        print(f"  [AVISO] Erro ao baixar: {ex} | URL: {url_dl[:80]}")
                if not downloaded:
                    print(f"  [ERRO] Não foi possível baixar o CSV de alunos — verifique URL_ALUNOS_HUB")
            except Exception as e:
                print(f"  [ERRO] Não foi possível baixar CSV de alunos: {e}")
        else:
            print(f"  [INFO] Relatorio_alunos_por_hub.csv não encontrado (defina URL_ALUNOS_HUB)")
    return p1, p2, tmpl, p3, p4, p5


def ler_excel(path, **kwargs):
    for engine in ('openpyxl', 'xlrd', None):
        try:
            kw = dict(kwargs)
            if engine: kw['engine'] = engine
            return pd.read_excel(path, **kw)
        except Exception: continue
    raise ValueError(f"Não foi possível ler {path} com nenhum engine disponível")


def processar(p1, p2):
    print(f"[{ts()}] Lendo tutores...")
    with open(p1, 'rb') as _f:
        _magic = _f.read(8); _preview = _f.read(200)
    print(f"  [DEBUG] Magic bytes de {p1}: {_magic.hex()}")
    if _magic[:2] == b'PK': print(f"  [DEBUG] Formato ZIP/XLSX confirmado")
    elif _magic[:2] in (b'\xd0\xcf', b'\xCF\xD0'): print(f"  [DEBUG] Formato XLS (OLE2) confirmado")
    else:
        print(f"  [ERRO] Arquivo não é Excel válido. Conteúdo inicial:")
        print((_magic + _preview).decode('utf-8', errors='replace')[:300])
        raise ValueError(f"Arquivo {p1} não é um Excel válido — SharePoint pode ter retornado HTML de erro")
    df_t = ler_excel(p1, sheet_name='Base de Tutores', header=1)
    col_sit  = next((c for c in df_t.columns if 'SITUA' in str(c).upper()), None)
    col_nome = next((c for c in df_t.columns if 'NOME'  in str(c).upper() and 'TUTOR' in str(c).upper()), None)
    col_polo = next((c for c in df_t.columns if c == 'POLO'), 'POLO')
    col_cur  = next((c for c in df_t.columns if c == 'CURSOS'), 'CURSOS')
    col_cat  = next((c for c in df_t.columns if 'CATEGORIA' in str(c).upper()), None)
    # PATCH 1: Detectar coluna CH SEMANAL na planilha de controle
    col_ch = next((c for c in df_t.columns if str(c).upper().strip() == 'CH SEMANAL' or
                   ('CH' in str(c).upper() and 'SEMAL' in str(c).upper())), None)
    if col_ch:
        print(f"[{ts()}] CH SEMANAL encontrada: '{col_ch}'")
        ch_vals = df_t[col_ch].dropna()
        print(f"[{ts()}] Amostra CH SEMANAL: {list(ch_vals.head(5))}")
    else:
        print(f"[{ts()}] CH SEMANAL não encontrada — colunas CH disponíveis: {[c for c in df_t.columns if 'CH' in str(c).upper()]}")
    df_at = df_t[df_t[col_sit].astype(str).str.strip() == 'Ativo'].copy() if col_sit else df_t.copy()
    df_at['_CHAVE'] = df_at[col_polo].astype(str).str.strip() + df_at[col_cur].astype(str).str.strip()
    print(f"[{ts()}] Lendo portfolios...")
    df_p = ler_excel(p2, sheet_name='Sheet1')
    def col(df, *partes):
        for c in df.columns:
            cu = str(c).upper()
            if all(p.upper() in cu for p in partes): return c
        return None
    c_chave = col(df_p, 'CHAVE', 'LINK')
    c_proto = col(df_p, 'PROTOCOLOS', 'ATIVIDADES')
    for sfx in (':7', ':8', ':6', ':9'):
        proto_sfx = [c for c in df_p.columns if 'PROTOCOLOS' in str(c).upper() and str(c).endswith(sfx)]
        if proto_sfx: c_proto = proto_sfx[0]; break
    else:
        proto_any = [c for c in df_p.columns if 'PROTOCOLOS' in str(c).upper() or 'ATIVIDADE' in str(c).upper()]
        if proto_any: c_proto = proto_any[0]
    data_cols = [c for c in df_p.columns if 'DATA DA APLICA' in str(c).upper() and str(c).endswith(':7')]
    if not data_cols: data_cols = [c for c in df_p.columns if 'DATA DA APLICA' in str(c).upper()]
    if not data_cols: data_cols = [c for c in df_p.columns if 'DATA' in str(c).upper() and 'APLICA' in str(c).upper()]
    c_data = data_cols[0] if data_cols else None
    def find_aluno_col(df):
        cols = df.columns.tolist()
        for suffix_end in ('72', '73', '74', '70', '71', '75'):
            for c in cols:
                if 'ESTUDANTES' in str(c).upper() and str(c).endswith(suffix_end): return c
        for c in cols:
            if 'ESTUDANTES' in str(c).upper() and 'PONTOS' not in str(c).upper(): return c
        for c in cols:
            cu = str(c).upper()
            if ('ALUNO' in cu or 'ALUNOS' in cu) and 'PONTOS' not in cu and 'COMENT' not in cu: return c
        for c in cols:
            cu = str(c).upper()
            if 'QUANTIDADE' in cu or 'QTD' in cu: return c
        return None
    c_aluno = find_aluno_col(df_p)
    cat_cols = [c for c in df_p.columns if 'CATEGORIA' in str(c).upper() and 'PONTOS' not in str(c).upper() and 'COMENT' not in str(c).upper()]
    c_cat = cat_cols[0] if cat_cols else None
    print(f"[{ts()}] Colunas: chave={c_chave}, proto={c_proto}, data={c_data}, alunos={c_aluno}, cat={c_cat}")
    c_ordem_cols = [c for c in df_p.columns if 'ORDEM' in str(c).upper() and 'PONTOS' not in str(c).upper() and 'COMENT' not in str(c).upper()]
    c_ordem = c_ordem_cols[0] if c_ordem_cols else None
    print(f"[{ts()}] Coluna ordem: {c_ordem}")
    df_p['_CHAVE']  = df_p[c_chave].astype(str).str.strip() if c_chave else ''
    df_p['_PROTO']  = df_p[c_proto].astype(str).str.strip() if c_proto else ''
    df_p['_DATA']   = pd.to_datetime(df_p[c_data], errors='coerce') if c_data else pd.NaT
    df_p['_ALUNOS'] = pd.to_numeric(df_p[c_aluno], errors='coerce').fillna(0).astype(int) if c_aluno else 0
    df_p['_CAT']    = df_p[c_cat].astype(str).str.strip() if c_cat else ''
    df_p['_ORDEM']  = df_p[c_ordem].astype(str).str.strip() if c_ordem else 'Ordem 1'
    catalogo_oficial = {}
    cat_file = os.path.join(SCRIPT_DIR, 'catalogo_oficial.json')
    if os.path.isfile(cat_file):
        with open(cat_file, encoding='utf-8') as f: raw = json.load(f)
        for cat_nome, praticas in raw.items():
            if isinstance(praticas, list) and praticas:
                if isinstance(praticas[0], dict): catalogo_oficial[cat_nome] = sorted(set(p['nome'] for p in praticas))
                else: catalogo_oficial[cat_nome] = sorted(set(praticas))
        print(f"[{ts()}] Catalogo oficial (JSON): {len(catalogo_oficial)} categorias")
    if not catalogo_oficial:
        cat_xlsx = achar_arquivo(SCRIPT_DIR, 'CATALOGO_EXPERIMENTOS.xlsx')
        if not cat_xlsx:
            for f in os.listdir(os.path.join(SCRIPT_DIR, 'planilhas')) if os.path.isdir(os.path.join(SCRIPT_DIR, 'planilhas')) else []:
                fu = f.upper()
                if ('RELAT' in fu and 'EXPER' in fu) or ('CATALOGO' in fu and 'EXPER' in fu):
                    cat_xlsx = os.path.join(SCRIPT_DIR, 'planilhas', f); break
        if cat_xlsx and os.path.isfile(cat_xlsx):
            try:
                df_cat = pd.read_excel(cat_xlsx)
                c_cat_nome = next((c for c in df_cat.columns if 'CATEGORIA' in str(c).upper()), None)
                c_exp_nome = next((c for c in df_cat.columns if 'EXPERIMENTO' in str(c).upper() or 'NOME' in str(c).upper()), None)
                c_sit = next((c for c in df_cat.columns if 'SITUA' in str(c).upper()), None)
                if c_cat_nome and c_exp_nome:
                    if c_sit: df_cat = df_cat[df_cat[c_sit].astype(str).str.strip().str.upper() == 'ATIVO']
                    for cat_val, grp in df_cat.groupby(c_cat_nome):
                        cat_str = str(cat_val).strip()
                        if cat_str and cat_str != 'nan':
                            nomes = sorted(set(str(n).strip() for n in grp[c_exp_nome].dropna() if str(n).strip() and str(n).strip() != 'nan'))
                            if nomes: catalogo_oficial[cat_str] = nomes
                    print(f"[{ts()}] Catalogo oficial (Excel): {len(catalogo_oficial)} categorias, {sum(len(v) for v in catalogo_oficial.values())} práticas")
            except Exception as e: print(f"[{ts()}] AVISO: Erro ao ler catálogo Excel: {e}")
    chave_to_cat_raw = {}; chave_to_cf = {}; chave_alias = {}
    polo_biofar_cursos = {}
    for _, t in df_at.iterrows():
        polo_   = str(t.get(col_polo, '') or '').strip()
        cursos_ = str(t.get(col_cur,  '') or '').strip()
        cat_    = str(t.get(col_cat,  '') or '').strip() if col_cat else ''
        if cursos_ in ('BBI', 'BFR') and 'BIO-FAR' in cat_.upper():
            if polo_ not in polo_biofar_cursos: polo_biofar_cursos[polo_] = set()
            polo_biofar_cursos[polo_].add(cursos_)
    for _, t in df_at.iterrows():
        polo = str(t.get(col_polo, '') or '').strip()
        cursos = str(t.get(col_cur, '') or '').strip()
        cat_raw = str(t.get(col_cat, '') or '').strip() if col_cat else ''
        cf = CAT_MAP.get(cat_raw, cat_raw)
        chave = polo + cursos
        if chave and cat_raw:
            chave_to_cat_raw[chave] = cat_raw
            chave_to_cf[chave] = cf
            if cursos in ('BBI', 'BFR'):
                outros = polo_biofar_cursos.get(polo, set()) - {cursos}
                variantes = [polo + 'BFI']
                for outro in outros:
                    variantes += [polo+cursos+'-'+outro, polo+outro+'-'+cursos, polo+cursos+outro, polo+outro+cursos]
                for v in variantes:
                    chave_to_cf.setdefault(v, cf)
                    chave_alias.setdefault(v, chave)
    oficial_p_to_cat = {}
    for cat, pracs in catalogo_oficial.items():
        for p in pracs: oficial_p_to_cat.setdefault(p, cat)
    catalogo_real = defaultdict(set)
    for _, r in df_p.iterrows():
        chave = str(r.get('_CHAVE', '') or '').strip()
        chave = chave_alias.get(chave, chave)
        proto = r['_PROTO']
        if not chave or chave == 'nan' or not proto or proto == 'nan': continue
        cf = chave_to_cf.get(chave, '')
        if not cf: continue
        for p in proto.split(';'):
            p = p.strip()
            if not p: continue
            cat_oficial = oficial_p_to_cat.get(p)
            if cat_oficial and cat_oficial != cf: continue
            catalogo_real[cf].add(p)
    catalogo = {}
    all_cats = set(list(catalogo_oficial.keys()) + list(catalogo_real.keys()))
    for cat in all_cats:
        base = set(catalogo_oficial.get(cat, [])); real = catalogo_real.get(cat, set())
        catalogo[cat] = sorted(base | real)
    print(f"[{ts()}] Catalogo final: {len(catalogo)} cats, {sum(len(v) for v in catalogo.values())} praticas")
    email_to_cf = {}; email_to_chave_tutor = {}
    col_email_t = next((c for c in df_t.columns if 'E-MAIL' in str(c).upper() or 'EMAIL' in str(c).upper()), None)
    if col_email_t:
        for _, t in df_at.iterrows():
            em = str(t.get(col_email_t, '') or '').strip().lower()
            chave_t = t['_CHAVE']
            cat_raw_ = str(t.get(col_cat, '') or '').strip() if col_cat else ''
            cf_ = CAT_MAP.get(cat_raw_, cat_raw_)
            if em and em != 'nan':
                email_to_cf[em] = cf_; email_to_chave_tutor[em] = chave_t
    col_email_p = next((c for c in df_p.columns if c.upper() in ('EMAIL', 'E-MAIL')), None)
    enviados = defaultdict(list)
    for _, r in df_p.iterrows():
        chave = r['_CHAVE']; proto = r['_PROTO']
        if not chave or chave == 'nan' or not proto or proto == 'nan': continue
        chave = chave_alias.get(chave, chave)
        if chave not in chave_to_cf and col_email_p:
            em_p = str(r.get(col_email_p, '') or '').strip().lower()
            if em_p in email_to_chave_tutor: chave = email_to_chave_tutor[em_p]
        data = r['_DATA']; aluno = int(r['_ALUNOS'])
        for p in proto.split(';'):
            p = p.strip()
            if p:
                ordem_val = str(r.get('_ORDEM', 'Ordem 1') or 'Ordem 1').strip()
                if not any(o in ordem_val for o in ['Ordem 1','Ordem 2','Ordem 3','Ordem 4','Ordem 5']):
                    ordem_val = 'Ordem 1'
                enviados[chave].append({'p': p[:80], 'd': str(data)[:10] if pd.notna(data) else None, 'a': aluno, 'o': ordem_val})
    tutores = []
    for _, t in df_at.iterrows():
        chave    = t['_CHAVE']
        cat_raw  = str(t.get(col_cat, '') or '').strip() if col_cat else ''
        cat_form = CAT_MAP.get(cat_raw, cat_raw)
        praticas = catalogo.get(cat_form, catalogo.get(cat_raw, []))
        hist     = enviados.get(chave, [])
        reais    = set(h['p'] for h in hist)
        pend     = [p for p in praticas if p not in reais]
        te = len(reais); tp = len(praticas)
        tutores.append({
            'n': str(t.get(col_nome, '') or ''),
            'p': str(t.get(col_polo, '') or ''),
            'c': cat_raw, 'cf': cat_form or 'Sem mapeamento',
            'tp': tp, 'te': te,
            'pend': pend, 'real': sorted(reais), 'hist': hist,
            'pct': round(te / tp * 100, 1) if tp else 0,
            'ch_semanal': _parse_ch(t.get(col_ch)) if col_ch else None,  # PATCH 1
        })
    seen = {}; tutores_dedup = []
    for t in tutores:
        key = (t.get('p',''), t.get('n','').strip().lower())
        if key in seen:
            existing = seen[key]
            existing['te'] = max(existing['te'], t['te'])
            existing['hist'] = existing['hist'] + [h for h in t['hist'] if h not in existing['hist']]
            existing['real'] = sorted(set(existing['real']) | set(t['real']))
            existing['pend'] = [p for p in existing['pend'] if p not in existing['real']]
            existing['tp'] = max(existing['tp'], t['tp'])
            existing['pct'] = round(existing['te'] / existing['tp'] * 100, 1) if existing['tp'] else 0
            if t.get('ch_semanal') and not existing.get('ch_semanal'):
                existing['ch_semanal'] = t['ch_semanal']
        else:
            seen[key] = t; tutores_dedup.append(t)
    tutores = tutores_dedup
    p_to_cat = {}
    for cat, pracs in catalogo.items():
        for p in pracs:
            if p not in p_to_cat: p_to_cat[p] = cat
    ps = defaultdict(lambda: {'enviou': 0, 'nao_enviou': 0, 'categoria': ''})
    for t in tutores:
        for p in t['real']:  ps[p]['enviou']    += 1
        for p in t['pend']:  ps[p]['nao_enviou'] += 1
    _p_fallback = {}
    for t in tutores:
        for p in t['real'] + t['pend']: _p_fallback.setdefault(p, t['cf'])
    for p in ps: ps[p]['categoria'] = p_to_cat.get(p, _p_fallback.get(p, ''))
    ps_all = sorted([{'nome': k, **v} for k, v in ps.items()], key=lambda x: -x['nao_enviou'])
    ps_list = ps_all[:30]
    cs = defaultdict(lambda: {'total_tutores': 0, 'com_100pct': 0, 'total_previstas': 0, 'total_enviadas': 0})
    for t in tutores:
        if not t['tp']: continue
        c = t['cf']
        cs[c]['total_tutores'] += 1
        if t['pct'] == 100: cs[c]['com_100pct'] += 1
        cs[c]['total_previstas'] += t['tp']; cs[c]['total_enviadas'] += t['te']
    print(f"[{ts()}] {len(tutores)} tutores, {sum(len(v) for v in catalogo.values())} praticas")
    prazos = PRAZOS_ORDENS.copy()
    hoje = datetime.now()
    status_ordem = {}
    # PATCH 3: usar datas de início reais de PERIODOS_ORDENS
    for ordem, prazo_str in prazos.items():
        prazo_date = datetime.strptime(prazo_str, '%d/%m/%Y')
        periodo = PERIODOS_ORDENS.get(ordem, {})
        inicio_str = periodo.get('inicio', '')
        if inicio_str:
            try: inicio_date = datetime.strptime(inicio_str, '%d/%m/%Y')
            except ValueError: inicio_date = prazo_date.replace(day=1)
        else: inicio_date = prazo_date.replace(day=1)
        if hoje > prazo_date: status_ordem[ordem] = 'VENCIDO'
        elif hoje >= inicio_date: status_ordem[ordem] = 'ABERTA'
        else: status_ordem[ordem] = 'FUTURA'
    tutores_out = []
    for t in tutores:
        por_ordem = {}
        for h in t['hist']:
            o = h.get('o', 'Ordem 1') or 'Ordem 1'
            por_ordem[o] = por_ordem.get(o, 0) + 1
        # PATCH 4: situação corrigida quando não há ordens vencidas
        ordens_vencidas = [o for o, s in status_ordem.items() if s == 'VENCIDO']
        ordens_abertas  = [o for o, s in status_ordem.items() if s == 'ABERTA']
        if not ordens_vencidas:
            if any(por_ordem.get(o, 0) > 0 for o in ordens_abertas) if ordens_abertas else False:
                sit = 'ok'
            elif ordens_abertas: sit = 'atrasado'
            else: sit = 'ok'
        else:
            if all(por_ordem.get(o, 0) > 0 for o in ordens_vencidas): sit = 'ok'
            elif any(por_ordem.get(o, 0) > 0 for o in ordens_vencidas): sit = 'atrasado'
            else: sit = 'urgente'
        tutores_out.append({
            **t,
            'nome': t.get('n',''), 'polo': t.get('p',''), 'cat': t.get('c',''),
            'n': t.get('n',''), 'p': t.get('p',''), 'c': t.get('c',''),
            'cf': t.get('cf','Sem mapeamento'),
            'por_ordem': por_ordem, 'porOrdem': por_ordem, 'situacao': sit,
        })
    col_email_key = next((c for c in df_t.columns if 'E-MAIL' in str(c).upper()), None)
    nome_to_email = {}
    if col_email_key:
        for _, row in df_at.iterrows():
            nome = str(row.get(col_nome, '') or '').strip()
            email = str(row.get(col_email_key, '') or '').strip().lower()
            if nome and email and email != 'nan': nome_to_email[nome] = email
    seen = {}; tutores_dedup = []
    for t in tutores_out:
        nome = t['n']; polo = t['p']
        email = nome_to_email.get(nome, '')
        key = email if email else (nome + '|' + polo).lower()
        if key not in seen:
            seen[key] = len(tutores_dedup); tutores_dedup.append(dict(t))
        else:
            ex = tutores_dedup[seen[key]]
            ex['hist'] = ex['hist'] + t['hist']
            merged_real = sorted(set(ex['real']) | set(t['real']))
            ex['real'] = merged_real; ex['te'] = len(merged_real)
            for o, cnt in t['por_ordem'].items():
                ex['por_ordem'][o] = ex['por_ordem'].get(o, 0) + cnt
                ex['porOrdem'][o]  = ex['porOrdem'].get(o, 0) + cnt
            real_set = set(merged_real)
            ex['pend'] = [p for p in ex['pend'] if p not in real_set]
            if ex['tp'] > 0: ex['pct'] = round(ex['te'] / ex['tp'] * 100, 1)
            # PATCH 5: reavalia situação + sincroniza sit
            _orv = [o for o, s in status_ordem.items() if s == 'VENCIDO']
            if not _orv: ex['situacao'] = 'ok'
            elif all(ex['por_ordem'].get(o, 0) > 0 for o in _orv): ex['situacao'] = 'ok'
            elif any(ex['por_ordem'].get(o, 0) > 0 for o in _orv): ex['situacao'] = 'atrasado'
            else: ex['situacao'] = 'urgente'
            ex['sit'] = ex['situacao']  # PATCH 5: sync shortcut
            if t.get('ch_semanal') and not ex.get('ch_semanal'): ex['ch_semanal'] = t['ch_semanal']
    tutores_out = tutores_dedup
    print(f"[{ts()}] Após deduplicação: {len(tutores_out)} tutores únicos")
    total     = len(tutores_out)
    enviaram  = sum(1 for t in tutores_out if t['te'] > 0)
    atrasados = sum(1 for t in tutores_out if t['situacao'] == 'atrasado')
    urgentes  = sum(1 for t in tutores_out if t['situacao'] == 'urgente')
    total_alunos = sum(h['a'] for t in tutores_out for h in t['hist'])
    polo_map = {}
    for t in tutores_out:
        p = t['polo']
        if p not in polo_map:
            polo_map[p] = {'POLO': p, 'polo': p, 'n': p, 'total': 0, 'enviaram': 0, 'atrasados': 0, 'alunos': 0}
        polo_map[p]['total'] += 1
        if t['te'] > 0: polo_map[p]['enviaram'] += 1
        if t['situacao'] == 'atrasado': polo_map[p]['atrasados'] += 1
        polo_map[p]['alunos'] += sum(h['a'] for h in t['hist'])
    polo_envios = {}
    for t in tutores_out:
        p = t['polo']
        polo_envios[p] = polo_envios.get(p, 0) + len(t.get('hist', []))
    polo_stats = sorted(polo_map.values(), key=lambda x: -x['atrasados'])
    for p in polo_stats:
        p['n'] = p.get('polo', p.get('POLO', ''))
        p['t'] = p['total']; p['e'] = p['enviaram']; p['a'] = p['alunos']
        p['pend'] = p['total'] - p['enviaram']
        p['pct'] = round(p['enviaram'] / p['total'] * 100) if p['total'] else 0
        p['envios'] = polo_envios.get(p['POLO'], 0)
    ordem_map = {o: {'envios': 0, 'alunos': 0} for o in prazos}
    for t in tutores_out:
        for h in t['hist']:
            o = h.get('o', 'Ordem 1') or 'Ordem 1'
            if o in ordem_map: ordem_map[o]['envios'] += 1; ordem_map[o]['alunos'] += h['a']
    por_ordem = [
        {'ordem': o, 'prazo': prazos[o], 'status': status_ordem[o],
         'envios': ordem_map[o]['envios'], 'alunos': ordem_map[o]['alunos']}
        for o in prazos
    ]
    mes_map = {}
    for t in tutores_out:
        for h in t['hist']:
            d = h.get('d') or ''
            mes = d[:7] if d and len(d) >= 7 else 'Sem data'
            if mes not in mes_map: mes_map[mes] = {'MES': mes, 'mes': mes, 'envios': 0, 'alunos': 0}
            mes_map[mes]['envios'] += 1; mes_map[mes]['alunos'] += h.get('a', 0)
    por_mes = sorted(mes_map.values(), key=lambda x: x['mes'])
    por_ordem_dict = {o: ordem_map[o]['envios'] for o in prazos}
    alunos_por_ordem = {o: ordem_map[o]['alunos'] for o in prazos}
    for p in polo_stats:
        p['n'] = p.get('polo', p.get('POLO', ''))
        p['t'] = p.get('total', 0); p['e'] = p.get('enviaram', 0); p['a'] = p.get('alunos', 0)
    for t in tutores_out:
        t['sit'] = t.get('situacao', 'urgente')
        t['al'] = sum(h.get('a', 0) for h in t.get('hist', []))
        t['email'] = nome_to_email.get(t.get('n', ''), '')
    praticas_template = []
    for p in ps_all:
        total_p = p['enviou'] + p['nao_enviou']
        praticas_template.append({
            'n': p['nome'], 'c': p['categoria'],
            'env_n': p['enviou'], 'pend_n': p['nao_enviou'],
            'pct': round(p['enviou'] / total_p * 100, 1) if total_p else 0,
            'nome': p['nome'], 'enviou': p['enviou'], 'nao_enviou': p['nao_enviou'], 'categoria': p['categoria'],
        })
    BRT = timezone(timedelta(hours=-3))
    gerado = datetime.now(BRT).strftime('%d/%m/%Y %H:%M')
    ch_ok = sum(1 for t in tutores_out if t.get('ch_semanal'))
    print(f"[{ts()}] {total} tutores · {enviaram} enviaram · {atrasados} atrasados · {urgentes} urgentes")
    print(f"[{ts()}] CH SEMANAL preenchida: {ch_ok}/{total} tutores")
    return limpar({
        'kpis': {
            'total': total, 'enviaram': enviaram, 'pendentes': total - enviaram,
            'atrasados': atrasados, 'urgentes': urgentes,
            'total_alunos': total_alunos, 'total_polos': len(polo_map),
            'polos_ok': sum(1 for p in polo_stats if p['enviaram'] > 0),
        },
        'tutores': tutores_out, 'polo_stats': polo_stats,
        'por_ordem': por_ordem_dict, 'por_ordem_lista': por_ordem,
        'alunos_por_ordem': alunos_por_ordem, 'status_ordem': status_ordem,
        'cat_stats': [{'categoria': k, **v} for k, v in cs.items()],
        'pratica_stats': ps_list, 'praticas': praticas_template,
        'catalogo': catalogo, 'prazos': prazos,
        'por_mes': por_mes, 'gerado_em': gerado,
    })


def _detectar_e_corrigir_base64(p4):
    import base64 as _b64
    try:
        with open(str(p4), 'rb') as f: raw = f.read(16)
        if raw[:4] == b'PK\x03\x04': return
        with open(str(p4), 'rb') as f: full = f.read()
        for padded in [full.strip(), full.strip() + b'==']:
            try:
                decoded = _b64.b64decode(padded)
                if decoded[:4] == b'PK\x03\x04':
                    with open(str(p4), 'wb') as fw: fw.write(decoded)
                    print(f"  [FIX] Base64 detectado e corrigido ({len(full)}->{len(decoded)} bytes)")
                    return
            except Exception: continue
        if raw[:5] in (b'\r\n<!D', b'<!DOC', b'<html'):
            print(f"  [ERRO] Arquivo é uma página HTML (login Microsoft)")
        else:
            print(f"  [INFO] Arquivo não é ZIP nem base64 ({raw[:4].hex()})")
    except Exception as e: print(f"  [AVISO] Verificação base64 falhou: {e}")


def _ler_lotacao_xlsx(p4):
    from openpyxl import load_workbook as _lwb
    wb = _lwb(str(p4), read_only=True, data_only=True, keep_vba=False)
    ws = wb['Quadro Geral de Lotação'] if 'Quadro Geral de Lotação' in wb.sheetnames else list(wb.worksheets)[0]
    return list(ws.iter_rows(values_only=True))

def _ler_lotacao_xls(p4):
    import xlrd
    wb = xlrd.open_workbook(str(p4))
    try: ws = wb.sheet_by_name('Quadro Geral de Lotação')
    except xlrd.XLRDError: ws = wb.sheet_by_index(0)
    rows = []
    for i in range(ws.nrows):
        row = []
        for j in range(ws.ncols):
            cell = ws.cell(i, j)
            if cell.ctype == xlrd.XL_CELL_DATE:
                import xlrd.xldate
                row.append(xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode))
            else: row.append(cell.value if cell.ctype != xlrd.XL_CELL_EMPTY else None)
        rows.append(tuple(row))
    return rows

def _ler_lotacao_pandas(p4):
    p = str(p4)
    try: df = pd.read_excel(p, sheet_name='Quadro Geral de Lotação', header=None)
    except Exception: df = pd.read_excel(p, sheet_name=0, header=None)
    return [tuple(row) for row in df.fillna('').values.tolist()]

def carregar_lotacao(p4):
    fname = os.path.basename(str(p4))
    print(f"[{ts()}] Lendo lotação de tutores ({fname})...")
    _detectar_e_corrigir_base64(p4)
    _rows = None
    for estrategia, fn in [('openpyxl', _ler_lotacao_xlsx), ('xlrd', _ler_lotacao_xls), ('pandas', _ler_lotacao_pandas)]:
        try:
            _rows = fn(p4)
            print(f"[{ts()}] Lotação lida via {estrategia}: {len(_rows)} linhas")
            break
        except Exception as e: print(f"[{ts()}] Tentativa {estrategia}: {e}")
    if not _rows: raise RuntimeError(f"Não foi possível ler {fname}")
    lotacao = {}
    for r in _rows[2:]:
        if not r[8] or str(r[8]).strip() in ('', '-', 'None', 'nan'): continue
        nome_raw = str(r[8]).strip(); nome_lower = nome_raw.lower()
        try: total_al = int(float(str(r[26] or 0)))
        except: total_al = 0
        lotacao[nome_lower] = {
            'nome_oficial': nome_raw, 'perfil': str(r[13] or '').strip(),
            'cursos': str(r[0] or '').strip(),
            'ch_semanal': _parse_ch(r[14]),  # reutiliza helper
            'ch_ideal': _parse_ch(r[15]) or 0.0,
            'contratacao': str(r[7] or '').strip(),
            'polo_hub': str(r[4] or '').strip(),
            'categoria_gio': str(r[29] or '').strip(),
            'total_alunos': total_al,
        }
        # Indexar também por nome primeiro+último para match mais abrangente
        _parts = nome_lower.split()
        if len(_parts) >= 2:
            _nfl = _parts[0] + ' ' + _parts[-1]
            if _nfl not in lotacao:
                lotacao[_nfl] = lotacao[nome_lower]
    print(f"[{ts()}] Lotação: {len(lotacao)} tutores mapeados")
    return lotacao


CURSOS_NOMES = {
    'EMF-ISN': 'Enfermagem e Instrumentação Cirúrgica', 'EMF-ISN2': 'Enfermagem e Instrumentação Cirúrgica',
    'BFR': 'Farmácia', 'BBI': 'Biomedicina', 'BFI': 'Fisioterapia', 'BTO': 'T. Ocupacional',
    'COS-TIP': 'Estética e Cosmética', 'NTR': 'Nutrição', 'AGM': 'Agronomia',
    'BAU': 'Arquitetura e Urbanismo', 'ECE-ENM-ENS-ENG-EEA-GPI-CDE-OBR-SAN-TER-FSA-SLF-QUI': 'Engenharias e Licenciaturas',
    'BIOMEDICINA': 'Biomedicina', 'FARMÁCIA': 'Farmácia', 'FISIOTERAPIA': 'Fisioterapia',
    'TERAPIA OCUPACIONAL': 'T. Ocupacional', 'NUTRIÇÃO': 'Nutrição', 'AGRONOMIA': 'Agronomia',
    'ARQUITETURA E URBANISMO': 'Arquitetura e Urbanismo',
}

def enriquecer_tutores(dados, lotacao):
    tutores = dados.get('tutores', [])
    matched = 0
    LAB_PARA_CAT = {
        'ENFERMAGEM,INSTRUMENTAÇÃO CIRÚRGICA': 'Enfermagem e Instrumentação Cirúrgica',
        'ENFERMAGEM,INSTRUMENTAÇÃO CIRÚRGICA2': 'Enfermagem e Instrumentação Cirúrgica',
        'BIOMEDICINA': 'Biomedicina', 'FARMÁCIA': 'Farmácia', 'FISIOTERAPIA': 'Fisioterapia',
        'TERAPIA OCUPACIONAL': 'T. Ocupacional',
        'TECNOLOGIA EM ESTÉTICA E COSMÉTICA,ESTÉTICA E IMAGEM PESSOAL': 'Estética e Cosmética',
        'NUTRIÇÃO': 'Nutrição', 'AGRONOMIA': 'Agronomia', 'ARQUITETURA E URBANISMO': 'Arquitetura e Urbanismo',
        'CONSTRUÇÃO DE EDIFÍCIOS,ENGENHARIA CIVIL,ENGENHARIA ELÉTRICA,ENGENHARIA DE PRODUÇÃO,ENGENHARIA MECÂNICA,ENGENHARIA AMBIENTAL E SANITÁRIA,FORMAÇÃO PEDAGÓGICA EM FÍSICA,FÍSICA,GESTÃO DA PRODUÇÃO INDUSTRIAL,CONTROLE DE OBRAS,QUÍMICA,SANEAMENTO AMBIENTAL,SEGUNDA LICENCIATURA EM FÍSICA,TECNOLOGIA EM ENERGIAS RENOVÁVEIS': 'Engenharias e Licenciaturas',
        'ENGENHARIA CIVIL,ENGENHARIA ELÉTRICA,ENGENHARIA DE PRODUÇÃO,ENGENHARIA MECÂNICA': 'Engenharias (Civil/Elét./Prod./Mec.)',
        'EMF-ISN': 'Enfermagem e Instrumentação Cirúrgica', 'BFR': 'Farmácia',
        'BBI': 'Biomedicina', 'BFI': 'Fisioterapia', 'BTO': 'T. Ocupacional',
        'COS-TIP': 'Estética e Cosmética', 'NTR': 'Nutrição', 'AGM': 'Agronomia',
        'BAU': 'Arquitetura e Urbanismo',
    }
    polo_lab_seen = set(); alunos_por_lab_raw = {}
    for nome_lower, info in lotacao.items():
        cursos_raw = info.get('cursos', '').strip().upper()
        polo_hub   = info.get('polo_hub', '').strip()
        total_al   = info.get('total_alunos', 0)
        if not cursos_raw or not total_al: continue
        chave_polo_lab = f"{polo_hub}||{cursos_raw}"
        if chave_polo_lab in polo_lab_seen: continue
        polo_lab_seen.add(chave_polo_lab)
        sep = '+' if '+' in cursos_raw else ','
        componentes = sorted([c.strip() for c in cursos_raw.split(sep)])
        lab_key = (','.join(componentes) if any(len(c) > 8 for c in componentes) else '+'.join(componentes))
        alunos_por_lab_raw[lab_key] = alunos_por_lab_raw.get(lab_key, 0) + total_al
    def _norm_lab_key(k):
        sep = '+' if '+' in k else ','
        partes = sorted([p.strip().upper() for p in k.split(sep)])
        return (','.join(partes) if any(len(p) > 8 for p in partes) else '+'.join(partes))
    lab_cat_norm = {_norm_lab_key(k): v for k, v in LAB_PARA_CAT.items()}
    alunos_por_curso = []
    for lab_key, total in sorted(alunos_por_lab_raw.items(), key=lambda x: -x[1]):
        nome = lab_cat_norm.get(lab_key)
        if not nome:
            primeiro = lab_key.split(',')[0].split('+')[0].strip()
            nome = CURSOS_NOMES.get(primeiro, primeiro.title())
        alunos_por_curso.append({'sigla': lab_key, 'curso': nome, 'alunos': total})
    dados['alunos_por_curso'] = alunos_por_curso
    total_al_sum = sum(x['alunos'] for x in alunos_por_curso)
    print(f"[{ts()}] Alunos por lab: {len(alunos_por_curso)} labs, total {total_al_sum:,}")
    for t in tutores:
        nome_lower = str(t.get('n', '')).lower()
        info = lotacao.get(nome_lower)
        if not info:
            for k, v in lotacao.items():
                if nome_lower in k or k in nome_lower: info = v; break
        if info:
            t['perfil'] = info['perfil']; t['cursos'] = info['cursos']
            # Lotação tem prioridade sobre CH da planilha de controle
            if info.get('ch_semanal') and info['ch_semanal'] > 0:
                t['ch_semanal'] = info['ch_semanal']
            t['ch_ideal'] = info.get('ch_ideal', 0)
            t['contratacao_lot'] = info['contratacao']
            t['lab'] = info.get('cursos', '')  # curso da planilha de lotação (para Multi 3)
            t['polo_hub_lot'] = info.get('polo_hub', '')
            matched += 1
    print(f"[{ts()}] Enriquecimento: {matched}/{len(tutores)} tutores com perfil/CH")
    dados['tutores'] = tutores
    return dados


def processar_gerenciamento_csv(p5):
    """Processa CSV detalhado de gerenciamento (REL_DETALHADO.csv)."""
    import csv, re as _re
    print(f"[{ts()}] Lendo gerenciamento (CSV)...")
    for enc in ('utf-8-sig', 'utf-8', 'latin-1', 'cp1252'):
        try:
            with open(str(p5), 'r', encoding=enc, errors='replace') as f:
                rows = list(csv.reader(f, delimiter=';'))
            break
        except: continue
    header = rows[0]; data = rows[1:]
    col = {h.strip().upper(): i for i, h in enumerate(header)}
    def gc(name): return col.get(name.upper())
    ci_polo = gc('LABORATORIO'); ci_cat = gc('CATEGORIA'); ci_exp = gc('NOME_EXPERIMENTO')
    ci_tutor = gc('TUTOR'); ci_mat = gc('ALUNOS_MATRICULADOS'); ci_agend = gc('ALUNOS_AGENDADOS')
    ci_capa = gc('CAPACIDADE_TOTAL'); ci_ofe = gc('OFERTAS_CADASTRADAS')
    ci_situ = gc('SITU_OFERTA'); ci_dt_ag = gc('DT_GERENCIADA'); ci_hr_ag = gc('HR_GERENCIADA')
    def gv(row, ci, default=''):
        try: return str(row[ci]).strip() if ci is not None and ci < len(row) else default
        except: return default
    def gn(row, ci):
        try: return float(str(row[ci]).replace(',','.').strip()) if ci is not None and ci < len(row) and row[ci] else 0
        except: return 0
    print(f"[{ts()}] Gerenciamento CSV: {len(data)} linhas, {len(header)} colunas")
    def extrair_ordem_exp(val):
        m = _re.match(r'O\.(\d+):\s*(.*)', str(val or ''))
        if m: return f'Ordem {m.group(1)}', m.group(2).strip()
        return '', str(val or '').strip()
    registros = []
    for r in data:
        polo = gv(r, ci_polo)
        if not polo: continue
        cat = gv(r, ci_cat); exp = gv(r, ci_exp); tutor = gv(r, ci_tutor); situ = gv(r, ci_situ)
        ordem, pratica = extrair_ordem_exp(exp)
        mat = int(gn(r, ci_mat)); agend = int(gn(r, ci_agend))
        capa = int(gn(r, ci_capa)); ofe = int(gn(r, ci_ofe))
        dt_ag = gv(r, ci_dt_ag); hr_ag = gv(r, ci_hr_ag)
        dt_ag_iso = ''
        if dt_ag and '/' in dt_ag:
            try:
                parts = dt_ag.split('/')
                dt_ag_iso = f"{parts[2]}-{parts[1]}-{parts[0]}"
            except: pass
        registros.append({
            'polo': polo, 'categoria': cat, 'pratica': pratica, 'ordem': ordem,
            'tutor': tutor if bool(tutor) else '',
            'tem_tutor': bool(tutor), 'tem_agenda': bool(dt_ag_iso),
            'gerenciado': ofe > 0, 'situ': situ,
            'alunos_mat': mat, 'alunos_agend': agend, 'capacidade': capa,
            'dt_agenda_iso': dt_ag_iso, 'hr_agenda': hr_ag,
        })
    df_r = pd.DataFrame(registros)
    if df_r.empty:
        return {'ger_kpis': {}, 'ger_polo': [], 'ger_cat': [], 'ger_ordem': [],
                'ger_contratacao': [], 'ger_agendas': [], 'ger_ofertas': []}
    polo_cat_tem_tutor = (
        df_r[df_r['tem_tutor']].groupby(['polo','categoria']).size()
        .reset_index(name='_qt').assign(_tem=True)
        .set_index(['polo','categoria'])['_tem']
    )
    def _fix_tem_tutor(row):
        return polo_cat_tem_tutor.get((row['polo'], row['categoria']), row['tem_tutor'])
    df_r['tem_tutor'] = df_r.apply(_fix_tem_tutor, axis=1)
    total = len(df_r); com_tutor = int(df_r['tem_tutor'].sum()); sem_tutor = total - com_tutor
    gerenciadas = int(df_r['gerenciado'].sum()); com_agenda = int(df_r['tem_agenda'].sum())
    tot_mat = int(df_r['alunos_mat'].sum()); tot_agend = int(df_r['alunos_agend'].sum())
    tot_capa = int(df_r['capacidade'].sum())
    print(f"[{ts()}] Gerenciamento: {total} ofertas, {gerenciadas} ger., {sem_tutor} sem tutor")
    ger_kpis = {
        'total_ofertas': total, 'ofertas_gerenciadas': gerenciadas,
        'ofertas_nao_gerenciadas': total - gerenciadas,
        'pct_gerenciado': round(gerenciadas/total*100,1) if total else 0,
        'ofertas_com_tutor': com_tutor, 'ofertas_sem_tutor': sem_tutor,
        'pct_com_tutor': round(com_tutor/total*100,1) if total else 0,
        'ofertas_com_agenda': com_agenda, 'total_alunos_matriculados': tot_mat,
        'total_alunos_agendados': tot_agend, 'total_capacidade': tot_capa,
        'pct_ocupacao': round(tot_agend/tot_capa*100,1) if tot_capa else 0,
        'polos_total': df_r['polo'].nunique(),
        'polos_sem_tutor': int(df_r[~df_r['tem_tutor']].groupby('polo').ngroups),
    }
    ger_polo = []
    for polo, grp in df_r.groupby('polo'):
        tuts = list(grp[grp['tem_tutor']]['tutor'].dropna().unique())
        ger_polo.append({
            'polo': str(polo), 'total_ofertas': len(grp),
            'gerenciadas': int(grp['gerenciado'].sum()),
            'pct_gerenciado': round(grp['gerenciado'].sum()/len(grp)*100,1) if len(grp) else 0,
            'com_tutor': int(grp['tem_tutor'].sum()), 'sem_tutor': int((~grp['tem_tutor']).sum()),
            'com_agenda': int(grp['tem_agenda'].sum()),
            'alunos_matriculados': int(grp['alunos_mat'].sum()), 'alunos_agendados': int(grp['alunos_agend'].sum()),
            'capacidade': int(grp['capacidade'].sum()), 'tutores_unicos': [str(t) for t in tuts],
        })
    ger_polo.sort(key=lambda x: -x['sem_tutor'])
    ger_cat = []
    for cat, grp in df_r.groupby('categoria'):
        ger_cat.append({
            'categoria': str(cat), 'total_ofertas': len(grp),
            'gerenciadas': int(grp['gerenciado'].sum()),
            'pct_gerenciado': round(grp['gerenciado'].sum()/len(grp)*100,1) if len(grp) else 0,
            'com_tutor': int(grp['tem_tutor'].sum()), 'sem_tutor': int((~grp['tem_tutor']).sum()),
            'alunos_matriculados': int(grp['alunos_mat'].sum()), 'alunos_agendados': int(grp['alunos_agend'].sum()),
        })
    ger_cat.sort(key=lambda x: -x['total_ofertas'])
    ger_ordem = []
    ordem_sort = {'Ordem 1':1,'Ordem 2':2,'Ordem 3':3,'Ordem 4':4,'Ordem 5':5}
    for ordem in sorted(df_r['ordem'].unique(), key=lambda x: ordem_sort.get(x,9)):
        if not ordem: continue
        grp = df_r[df_r['ordem']==ordem]
        ger_ordem.append({
            'ordem': ordem, 'total_ofertas': len(grp),
            'gerenciadas': int(grp['gerenciado'].sum()),
            'pct_gerenciado': round(grp['gerenciado'].sum()/len(grp)*100,1) if len(grp) else 0,
            'com_tutor': int(grp['tem_tutor'].sum()),
            'alunos_matriculados': int(grp['alunos_mat'].sum()), 'alunos_agendados': int(grp['alunos_agend'].sum()),
            'dt_inicio': '', 'dt_fim': PRAZOS_ORDENS.get(ordem,''),
        })
    ger_contratacao = []
    for (polo, cat), grp in df_r.groupby(['polo','categoria']):
        tuts = list(grp[grp['tem_tutor']]['tutor'].dropna().unique())
        ger_contratacao.append({
            'polo': str(polo), 'categoria': str(cat), 'total_ofertas': len(grp),
            'tem_tutor': len(tuts)>0, 'tutores': [str(t) for t in tuts],
            'status': 'Contratado' if len(tuts)>0 else 'Sem tutor',
        })
    ger_agendas = []
    for polo, grp in df_r.groupby('polo'):
        total_p = len(grp); com_ag = int(grp['tem_agenda'].sum()); sem_ag = total_p - com_ag
        datas_por_cat = {}; datas_por_tutor = {}
        for _, row in grp[grp['tem_agenda']].iterrows():
            d = row['dt_agenda_iso']; c = row['categoria']; t = row['tutor'] or ''
            if d:
                if d not in datas_por_cat: datas_por_cat[d] = {'cats': [], 'tutores': []}
                if c and c not in datas_por_cat[d]['cats']: datas_por_cat[d]['cats'].append(c)
                if t and t not in datas_por_cat[d]['tutores']: datas_por_cat[d]['tutores'].append(t)
        # PATCH 7: estrutura completa de datas_por_tutor
        for d, v in datas_por_cat.items():
            datas_por_tutor[d] = v['tutores']
        ger_agendas.append({
            'polo': str(polo), 'total': total_p, 'com_agenda': com_ag, 'sem_agenda': sem_ag,
            'pct_agendado': round(com_ag/total_p*100, 1) if total_p else 0,
            'datas_agenda': sorted(datas_por_cat.keys()),
            'datas_por_cat': {d: v['cats'] for d, v in datas_por_cat.items()},
            'datas_por_tutor': datas_por_tutor,  # PATCH 7
        })
    ger_agendas.sort(key=lambda x: -x['sem_agenda'])
    ger_ofertas_detalhe = []
    for _, row in df_r.iterrows():
        ger_ofertas_detalhe.append({
            'polo': row['polo'], 'categoria': row['categoria'],
            'ordem': row['ordem'], 'pratica': row['pratica'],
            'tutor': row['tutor'], 'tem_tutor': row['tem_tutor'],
            'tem_agenda': row['tem_agenda'], 'gerenciado': row['gerenciado'],
            'alunos_mat': row['alunos_mat'], 'alunos_agend': row['alunos_agend'],
            'dt_agenda': row['dt_agenda_iso'], 'hr_agenda': row['hr_agenda'],
        })
    return {
        'ger_kpis': ger_kpis, 'ger_polo': ger_polo, 'ger_cat': ger_cat,
        'ger_ordem': ger_ordem, 'ger_contratacao': ger_contratacao,
        'ger_agendas': ger_agendas, 'ger_ofertas': ger_ofertas_detalhe,
    }


def _processar_gerenciamento_novo(df_g):
    import re as _re
    col = {str(c).strip().upper(): c for c in df_g.columns}
    def gc(name): return col.get(name.upper())
    c_polo = gc('LABORATORIO'); c_cat = gc('CATEGORIA'); c_exp = gc('NOME_EXPERIMENTO')
    c_tutor = gc('TUTOR'); c_mat = gc('ALUNOS_MATRICULADOS'); c_agend = gc('ALUNOS_AGENDADOS')
    c_capa = gc('CAPACIDADE_TOTAL'); c_ofe = gc('OFERTAS_CADASTRADAS'); c_situ = gc('SITU_OFERTA')
    c_dt_ag = gc('DT_GERENCIADA'); c_hr_ag = gc('HR_GERENCIADA')
    def extrair_ordem_exp(val):
        m = _re.match(r'O\.(\d+):\s*(.*)', str(val or ''))
        if m: return f'Ordem {m.group(1)}', m.group(2).strip()
        return '', str(val or '').strip()
    df = df_g.copy()
    df['_POLO']  = df[c_polo].astype(str).str.strip() if c_polo else ''
    df['_CAT']   = df[c_cat].astype(str).str.strip()  if c_cat  else ''
    df['_TUTOR'] = df[c_tutor].fillna('').astype(str).str.strip().replace('nan','') if c_tutor else ''
    df['_MAT']   = pd.to_numeric(df[c_mat],  errors='coerce').fillna(0).astype(int) if c_mat  else 0
    df['_AGEND'] = pd.to_numeric(df[c_agend],errors='coerce').fillna(0).astype(int) if c_agend else 0
    df['_CAPA']  = pd.to_numeric(df[c_capa], errors='coerce').fillna(0).astype(int) if c_capa  else 0
    df['_OFE']   = pd.to_numeric(df[c_ofe],  errors='coerce').fillna(0).astype(int) if c_ofe   else 0
    df['_TEM_TUTOR'] = df['_TUTOR'].str.len() > 0
    _situ_col = df[c_situ].fillna('').astype(str).str.strip() if c_situ else pd.Series([''] * len(df))
    # GERENCIADO = tem tutor E (tem ofertas cadastradas OU status CONCLUÍDO)
    df['_GERENCIADO'] = df['_TEM_TUTOR'] & ((df['_OFE'] > 0) | _situ_col.str.upper().str.contains('CONCLU', na=False))
    dt_col = df[c_dt_ag] if c_dt_ag else pd.Series([''] * len(df))
    def to_iso(v):
        if v is None: return ''
        try:
            import datetime as _dt
            if isinstance(v, (_dt.datetime, _dt.date)): return v.strftime('%Y-%m-%d')
        except: pass
        sv = str(v).strip()
        if not sv or sv == 'nan': return ''
        if '/' in sv:
            try:
                parts = sv.split('/')
                if len(parts) == 3: return f'{parts[2]}-{parts[1].zfill(2)}-{parts[0].zfill(2)}'
            except: pass
        if '-' in sv and len(sv) >= 10: return sv[:10]
        try:
            n = float(sv)
            import datetime as _dt
            base = _dt.date(1899, 12, 30)
            return (base + _dt.timedelta(days=int(n))).strftime('%Y-%m-%d')
        except: pass
        return ''
    df['_DT_AG_ISO'] = dt_col.apply(to_iso)
    df['_TEM_AGENDA'] = df['_DT_AG_ISO'].str.len() > 0
    df['_HR_AG'] = df[c_hr_ag].fillna('').astype(str).str.strip().replace('nan','').replace('NaT','') if c_hr_ag else ''
    parsed = (df[c_exp] if c_exp else pd.Series([''] * len(df))).apply(extrair_ordem_exp)
    df['_ORDEM'] = parsed.apply(lambda x: x[0])
    df['_PRATICA'] = parsed.apply(lambda x: x[1])
    df = df[df['_POLO'].str.len() > 0].copy()
    total = len(df); com_tutor = int(df['_TEM_TUTOR'].sum()); gerenciadas = int(df['_GERENCIADO'].sum())
    com_agenda = int(df['_TEM_AGENDA'].sum())
    # FIX: Alunos Matriculados — deduplicar por polo×categoria (remove contagem múltipla por ordem)
    _mat_col = '_MAT'; _agend_col = '_AGEND'; _capa_col = '_CAPA'
    _raw_mat = int(df[_mat_col].sum()) if _mat_col in df.columns else 0
    _grp_cols_ok = ['_POLO','_CAT']
    if all(c in df.columns for c in _grp_cols_ok + [_mat_col, _agend_col, _capa_col]):
        _dedup = df.groupby(_grp_cols_ok)[[_mat_col, _agend_col, _capa_col]].max()
        tot_mat   = int(_dedup[_mat_col].sum())
        tot_agend = int(_dedup[_agend_col].sum())
        tot_capa  = int(_dedup[_capa_col].sum())
        print(f"[{ts()}] Alunos DEDUPLICADOS por polo×cat: {tot_mat:,} (bruto era {_raw_mat:,}, redução: {_raw_mat-tot_mat:,})")
    else:
        tot_mat   = _raw_mat
        tot_agend = int(df[_agend_col].sum()) if _agend_col in df.columns else 0
        tot_capa  = int(df[_capa_col].sum())  if _capa_col  in df.columns else 0
        print(f"[{ts()}] Alunos sem dedup: {tot_mat:,}")
    print(f"[{ts()}] Gerenciamento: {total} ofertas, {gerenciadas} ger., {total-com_tutor} sem tutor")
    print(f"[{ts()}] Agendas: {com_agenda} · datas: {sorted(df[df['_TEM_AGENDA']]['_DT_AG_ISO'].head(3).tolist())}")
    print(f"[{ts()}] {df['_POLO'].nunique()} polos, {df['_CAT'].nunique()} cats, {df['_ORDEM'].nunique()} ordens")
    ger_kpis = {
        'total_ofertas': total, 'ofertas_gerenciadas': gerenciadas,
        'ofertas_nao_gerenciadas': total - gerenciadas,
        'pct_gerenciado': round(gerenciadas/total*100,1) if total else 0,
        'ofertas_com_tutor': com_tutor, 'ofertas_sem_tutor': total-com_tutor,
        'pct_com_tutor': round(com_tutor/total*100,1) if total else 0,
        'ofertas_com_agenda': com_agenda, 'total_alunos_matriculados': tot_mat,
        'total_alunos_agendados': tot_agend, 'total_capacidade': tot_capa,
        'pct_ocupacao': round(tot_agend/tot_capa*100,1) if tot_capa else 0,
        'polos_total': df['_POLO'].nunique(),
        'polos_sem_tutor': int(df[~df['_TEM_TUTOR']].groupby('_POLO').ngroups),
    }
    ger_polo = []
    for polo, grp in df.groupby('_POLO'):
        tuts = list(grp[grp['_TEM_TUTOR']]['_TUTOR'].dropna().unique())
        ger_polo.append({
            'polo': str(polo), 'total_ofertas': len(grp),
            'gerenciadas': int(grp['_GERENCIADO'].sum()),
            'pct_gerenciado': round(grp['_GERENCIADO'].sum()/len(grp)*100,1) if len(grp) else 0,
            'com_tutor': int(grp['_TEM_TUTOR'].sum()), 'sem_tutor': int((~grp['_TEM_TUTOR']).sum()),
            'com_agenda': int(grp['_TEM_AGENDA'].sum()),
            'alunos_matriculados': int(grp['_MAT'].sum()), 'alunos_agendados': int(grp['_AGEND'].sum()),
            'capacidade': int(grp['_CAPA'].sum()), 'tutores_unicos': [str(t) for t in tuts],
        })
    ger_polo.sort(key=lambda x: -x['sem_tutor'])
    ger_cat = []
    for cat, grp in df.groupby('_CAT'):
        ger_cat.append({
            'categoria': str(cat), 'total_ofertas': len(grp),
            'gerenciadas': int(grp['_GERENCIADO'].sum()),
            'pct_gerenciado': round(grp['_GERENCIADO'].sum()/len(grp)*100,1) if len(grp) else 0,
            'com_tutor': int(grp['_TEM_TUTOR'].sum()), 'sem_tutor': int((~grp['_TEM_TUTOR']).sum()),
            'alunos_matriculados': int(grp['_MAT'].sum()), 'alunos_agendados': int(grp['_AGEND'].sum()),
        })
    ger_cat.sort(key=lambda x: -x['total_ofertas'])
    ger_ordem = []; ordem_sort = {'Ordem 1':1,'Ordem 2':2,'Ordem 3':3,'Ordem 4':4,'Ordem 5':5}
    for ordem in sorted(df['_ORDEM'].unique(), key=lambda x: ordem_sort.get(x,9)):
        if not ordem: continue
        grp = df[df['_ORDEM']==ordem]
        ger_ordem.append({
            'ordem': ordem, 'total_ofertas': len(grp),
            'gerenciadas': int(grp['_GERENCIADO'].sum()),
            'pct_gerenciado': round(grp['_GERENCIADO'].sum()/len(grp)*100,1) if len(grp) else 0,
            'com_tutor': int(grp['_TEM_TUTOR'].sum()),
            'alunos_matriculados': int(grp['_MAT'].sum()), 'alunos_agendados': int(grp['_AGEND'].sum()),
            'dt_inicio': '', 'dt_fim': PRAZOS_ORDENS.get(ordem,''),
        })
    ger_contratacao = []
    for (polo, cat), grp in df.groupby(['_POLO','_CAT']):
        tuts = list(grp[grp['_TEM_TUTOR']]['_TUTOR'].dropna().unique())
        ger_contratacao.append({
            'polo': str(polo), 'categoria': str(cat), 'total_ofertas': len(grp),
            'tem_tutor': len(tuts)>0, 'tutores': [str(t) for t in tuts],
            'status': 'Contratado' if len(tuts)>0 else 'Sem tutor',
        })
    ger_agendas = []
    for polo, grp in df.groupby('_POLO'):
        total_p = len(grp); com_ag = int(grp['_TEM_AGENDA'].sum())
        datas_por_cat = {}; datas_por_tutor = {}
        for _, row in grp[grp['_TEM_AGENDA']].iterrows():
            d = row['_DT_AG_ISO']; c = row['_CAT']; t = row['_TUTOR']
            if d:
                if d not in datas_por_cat: datas_por_cat[d]=[]
                if c and c not in datas_por_cat[d]: datas_por_cat[d].append(c)
                if d not in datas_por_tutor: datas_por_tutor[d]=[]
                if t and t not in datas_por_tutor[d]: datas_por_tutor[d].append(t)
        ger_agendas.append({
            'polo': str(polo), 'total': total_p, 'com_agenda': com_ag,
            'sem_agenda': total_p-com_ag,
            'pct_agendado': round(com_ag/total_p*100,1) if total_p else 0,
            'datas_agenda': sorted(datas_por_cat.keys()),
            'datas_por_cat': datas_por_cat,
            'datas_por_tutor': datas_por_tutor,  # PATCH 7: preservado
        })
    ger_agendas.sort(key=lambda x: -x['sem_agenda'])
    ger_ofertas = []
    for _, row in df.iterrows():
        ger_ofertas.append({
            'polo': row['_POLO'], 'categoria': row['_CAT'],
            'ordem': row['_ORDEM'], 'pratica': row['_PRATICA'],
            'tutor': row['_TUTOR'], 'tem_tutor': bool(row['_TEM_TUTOR']),
            'tem_agenda': bool(row['_TEM_AGENDA']), 'gerenciado': bool(row['_GERENCIADO']),
            'alunos_mat': int(row['_MAT']), 'alunos_agend': int(row['_AGEND']),
            'dt_agenda': row['_DT_AG_ISO'], 'hr_agenda': row['_HR_AG'],
        })
    return {
        'ger_kpis': ger_kpis, 'ger_polo': ger_polo, 'ger_cat': ger_cat,
        'ger_ordem': ger_ordem, 'ger_contratacao': ger_contratacao,
        'ger_agendas': ger_agendas, 'ger_ofertas': ger_ofertas,
    }


def processar_gerenciamento(p3):
    print(f"[{ts()}] Lendo gerenciamento...")
    df_g = ler_excel(p3)
    print(f"[{ts()}] Gerenciamento: {len(df_g)} linhas, {len(df_g.columns)} colunas")
    cols_upper = [str(c).upper() for c in df_g.columns]
    is_novo = 'LABORATORIO' in cols_upper and 'NOME_EXPERIMENTO' in cols_upper
    if is_novo:
        print(f"[{ts()}] Formato: NOVO (relatório detalhado)")
        return _processar_gerenciamento_novo(df_g)
    print(f"[{ts()}] Formato: ANTIGO (GIOCONDA)")
    def gcol(df, *partes):
        for c in df.columns:
            cu = str(c).upper()
            if all(p.upper() in cu for p in partes): return c
        return None
    c_polo = gcol(df_g, 'CEEM', 'RSOC') or 'CEEM_RSOC'
    c_cat  = gcol(df_g, 'CATP', 'NOME') or 'CATP_NOME'
    c_lab  = gcol(df_g, 'LABE', 'NOME') or 'LABE_NOME'
    c_curso = gcol(df_g, 'NOME', 'CURS') or 'NOME_CURS'
    c_situ = gcol(df_g, 'SITU') or 'SITU'
    c_alunos = gcol(df_g, 'ALUNOS', 'MATRIC') or 'ALUNOS_MATRICULADOS'
    c_capa_exp = gcol(df_g, 'CAPA', 'EXP') or 'CAPA_EXP'
    c_ofe_cad = gcol(df_g, 'OFE', 'CAD') or 'OFE_CAD'
    c_qtd_alun = gcol(df_g, 'QTD', 'ALUN') or 'QTD_ALUN'
    c_tutor = gcol(df_g, 'TUTOR') or 'TUTOR'
    c_dt_agenda = gcol(df_g, 'DT', 'GERENCIADA') or 'DT_GERENCIADA'
    c_hr_agenda = gcol(df_g, 'HR', 'GERENCIADA') or 'HR_GERENCIADA'
    c_ofex_dtin = gcol(df_g, 'OFEX', 'DTIN') or 'OFEX_DTIN'
    c_ofex_dtfi = gcol(df_g, 'OFEX', 'DTFI') or 'OFEX_DTFI'
    if c_situ in df_g.columns:
        df_g = df_g[df_g[c_situ].astype(str).str.strip().str.upper() == 'ATIVO'].copy()
    print(f"[{ts()}] Gerenciamento após filtro ativos: {len(df_g)} linhas")
    df_g['_ORDEM_G'] = ''; df_g['_PRATICA_G'] = ''
    if c_lab in df_g.columns:
        import re
        def extrair_ordem(val):
            val = str(val or '')
            m = re.match(r'O\.(\d+):\s*(.*)', val)
            if m: return f'Ordem {m.group(1)}', m.group(2).strip()
            return '', val.strip()
        parsed = df_g[c_lab].apply(extrair_ordem)
        df_g['_ORDEM_G'] = parsed.apply(lambda x: x[0])
        df_g['_PRATICA_G'] = parsed.apply(lambda x: x[1])
    # FIX BUG 1: _TEM_TUTOR deve ser definido ANTES de _GERENCIADO
    df_g['_TEM_TUTOR'] = df_g[c_tutor].notna() & (df_g[c_tutor].astype(str).str.strip() != '') & (df_g[c_tutor].astype(str).str.strip().str.upper() != 'NAN')
    # GERENCIADO = tem tutor E (tem ofertas cadastradas OU status CONCLUÍDO)
    _situ_g = df_g[c_situ].fillna('').astype(str).str.strip() if c_situ and c_situ in df_g.columns else pd.Series([''] * len(df_g))
    df_g['_GERENCIADO'] = df_g['_TEM_TUTOR'] & ((pd.to_numeric(df_g.get(c_ofe_cad, 0), errors='coerce').fillna(0) > 0) | _situ_g.str.upper().str.contains('CONCLU', na=False))
    df_g['_TEM_AGENDA'] = df_g.get(c_dt_agenda, pd.Series(dtype='object')).notna()
    df_g['_ALUNOS_MAT'] = pd.to_numeric(df_g.get(c_alunos, 0), errors='coerce').fillna(0).astype(int)
    df_g['_QTD_ALUN'] = pd.to_numeric(df_g.get(c_qtd_alun, 0), errors='coerce').fillna(0).astype(int)
    df_g['_CAPA'] = pd.to_numeric(df_g.get(c_capa_exp, 0), errors='coerce').fillna(0).astype(int)
    total_ofertas = len(df_g); gerenciadas = int(df_g['_GERENCIADO'].sum())
    com_tutor = int(df_g['_TEM_TUTOR'].sum()); sem_tutor = total_ofertas - com_tutor
    # FIX: Alunos Matriculados — deduplicar por polo×categoria (soma bruta conta os mesmos alunos por ordem)
    # Usar apenas colunas que REALMENTE existem (não fallbacks)
    _c_polo_real = c_polo if (c_polo and c_polo in df_g.columns) else None
    _c_cat_real  = c_cat  if (c_cat  and c_cat  in df_g.columns) else None
    # Se nenhuma das buscas primárias funcionou, tentar qualquer coluna polo/cat
    if not _c_polo_real:
        _c_polo_real = next((c for c in df_g.columns if 'POLO' in str(c).upper() or 'CEEM' in str(c).upper()), None)
    if not _c_cat_real:
        _c_cat_real = next((c for c in df_g.columns if 'CATEG' in str(c).upper() or 'CATP' in str(c).upper()), None)
    _raw_mat = int(df_g['_ALUNOS_MAT'].sum())
    if _c_polo_real and _c_cat_real:
        _dedup_g = df_g.groupby([_c_polo_real, _c_cat_real])[['_ALUNOS_MAT','_QTD_ALUN','_CAPA']].max()
        tot_mat   = int(_dedup_g['_ALUNOS_MAT'].sum())
        tot_agend = int(_dedup_g['_QTD_ALUN'].sum())
        tot_capa  = int(_dedup_g['_CAPA'].sum())
        print(f"[{ts()}] Alunos DEDUPLICADOS por polo×cat: {tot_mat:,} (bruto era {_raw_mat:,}, redução: {_raw_mat-tot_mat:,})")
    else:
        tot_mat = _raw_mat
        tot_agend = int(df_g['_QTD_ALUN'].sum())
        tot_capa  = int(df_g['_CAPA'].sum())
        print(f"[{ts()}] Alunos sem dedup (colunas polo/cat não encontradas): {tot_mat:,}")
    polos_total = df_g[c_polo].nunique() if c_polo in df_g.columns else 0
    polos_sem_tutor_count = int(df_g[~df_g['_TEM_TUTOR']].groupby(c_polo).ngroups) if c_polo in df_g.columns else 0
    ger_kpis = {
        'total_ofertas': total_ofertas, 'ofertas_gerenciadas': gerenciadas,
        'ofertas_nao_gerenciadas': total_ofertas - gerenciadas,
        'pct_gerenciado': round(gerenciadas/total_ofertas*100,1) if total_ofertas else 0,
        'ofertas_com_tutor': com_tutor, 'ofertas_sem_tutor': sem_tutor,
        'pct_com_tutor': round(com_tutor/total_ofertas*100,1) if total_ofertas else 0,
        'ofertas_com_agenda': int(df_g['_TEM_AGENDA'].sum()),
        'total_alunos_matriculados': tot_mat, 'total_alunos_agendados': tot_agend,
        'total_capacidade': tot_capa,
        'pct_ocupacao': round(tot_agend/tot_capa*100,1) if tot_capa else 0,
        'polos_total': polos_total, 'polos_sem_tutor': polos_sem_tutor_count,
    }
    print(f"[{ts()}] Gerenciamento: {total_ofertas} ofertas, {gerenciadas} ger., {sem_tutor} sem tutor")
    ger_polo = []
    if c_polo in df_g.columns:
        for polo, grp in df_g.groupby(c_polo):
            ger_polo.append({
                'polo': str(polo), 'total_ofertas': len(grp),
                'gerenciadas': int(grp['_GERENCIADO'].sum()),
                'pct_gerenciado': round(grp['_GERENCIADO'].sum()/len(grp)*100,1) if len(grp) else 0,
                'com_tutor': int(grp['_TEM_TUTOR'].sum()), 'sem_tutor': int((~grp['_TEM_TUTOR']).sum()),
                'com_agenda': int(grp['_TEM_AGENDA'].sum()),
                'alunos_matriculados': int(grp['_ALUNOS_MAT'].sum()), 'alunos_agendados': int(grp['_QTD_ALUN'].sum()),
                'capacidade': int(grp['_CAPA'].sum()),
                'tutores_unicos': list(grp[grp['_TEM_TUTOR']][c_tutor].dropna().unique()),
            })
        ger_polo.sort(key=lambda x: -x['sem_tutor'])
    ger_cat = []
    if c_cat in df_g.columns:
        for cat, grp in df_g.groupby(c_cat):
            ger_cat.append({
                'categoria': str(cat), 'total_ofertas': len(grp),
                'gerenciadas': int(grp['_GERENCIADO'].sum()),
                'pct_gerenciado': round(grp['_GERENCIADO'].sum()/len(grp)*100,1) if len(grp) else 0,
                'com_tutor': int(grp['_TEM_TUTOR'].sum()), 'sem_tutor': int((~grp['_TEM_TUTOR']).sum()),
                'alunos_matriculados': int(grp['_ALUNOS_MAT'].sum()), 'alunos_agendados': int(grp['_QTD_ALUN'].sum()),
            })
        ger_cat.sort(key=lambda x: -x['total_ofertas'])
    ger_ordem = []
    ordens_validas = [o for o in sorted(df_g['_ORDEM_G'].unique()) if o and 'Ordem' in str(o)]
    for ordem in ordens_validas:
        grp = df_g[df_g['_ORDEM_G'] == ordem]
        datas_inicio = pd.to_datetime(grp.get(c_ofex_dtin, pd.Series(dtype='object')), errors='coerce').dropna()
        datas_fim = pd.to_datetime(grp.get(c_ofex_dtfi, pd.Series(dtype='object')), errors='coerce').dropna()
        dt_inicio = datas_inicio.min().strftime('%d/%m/%Y') if len(datas_inicio) > 0 else ''
        dt_fim = datas_fim.max().strftime('%d/%m/%Y') if len(datas_fim) > 0 else ''
        ger_ordem.append({
            'ordem': ordem, 'total_ofertas': len(grp),
            'gerenciadas': int(grp['_GERENCIADO'].sum()),
            'pct_gerenciado': round(grp['_GERENCIADO'].sum()/len(grp)*100,1) if len(grp) else 0,
            'com_tutor': int(grp['_TEM_TUTOR'].sum()),
            'alunos_matriculados': int(grp['_ALUNOS_MAT'].sum()), 'alunos_agendados': int(grp['_QTD_ALUN'].sum()),
            'dt_inicio': dt_inicio, 'dt_fim': dt_fim,
        })
    ger_contratacao = []
    if c_polo in df_g.columns and c_cat in df_g.columns:
        for (polo, cat), grp in df_g.groupby([c_polo, c_cat]):
            tutores_list = list(grp[grp['_TEM_TUTOR']][c_tutor].dropna().unique())
            ger_contratacao.append({
                'polo': str(polo), 'categoria': str(cat), 'total_ofertas': len(grp),
                'tem_tutor': len(tutores_list)>0, 'tutores': [str(t) for t in tutores_list],
                'status': 'Contratado' if len(tutores_list)>0 else 'Sem tutor',
            })
        ger_contratacao.sort(key=lambda x: (0 if x['tem_tutor'] else 1, x['polo']))
    ger_agendas = []
    if c_polo in df_g.columns:
        for polo, grp in df_g.groupby(c_polo):
            total = len(grp); com_agenda = int(grp['_TEM_AGENDA'].sum()); sem_agenda = total - com_agenda
            datas = []; datas_por_cat = {}; datas_por_tutor = {}
            if c_dt_agenda and c_dt_agenda in grp.columns:
                for _, ag_row in grp[grp['_TEM_AGENDA']].iterrows():
                    dt_val = pd.to_datetime(ag_row.get(c_dt_agenda), errors='coerce')
                    if pd.notna(dt_val):
                        dt_str = dt_val.strftime('%Y-%m-%d')
                        cat_val = str(ag_row.get(c_cat, '') or '')
                        tutor_val = str(ag_row.get(c_tutor, '') or '')
                        if dt_str not in datas: datas.append(dt_str)
                        if cat_val:
                            if dt_str not in datas_por_cat: datas_por_cat[dt_str] = []
                            if cat_val not in datas_por_cat[dt_str]: datas_por_cat[dt_str].append(cat_val)
                        if tutor_val and tutor_val != 'nan':
                            if dt_str not in datas_por_tutor: datas_por_tutor[dt_str] = []
                            if tutor_val not in datas_por_tutor[dt_str]: datas_por_tutor[dt_str].append(tutor_val)
                datas = sorted(set(datas))
            ger_agendas.append({
                'polo': str(polo), 'total': total, 'com_agenda': com_agenda, 'sem_agenda': sem_agenda,
                'pct_agendado': round(com_agenda/total*100, 1) if total else 0,
                'datas_agenda': datas, 'datas_por_cat': datas_por_cat,
                'datas_por_tutor': datas_por_tutor,  # PATCH 7: preservado
            })
        ger_agendas.sort(key=lambda x: -x['sem_agenda'])
    ger_ofertas_detalhe = []
    for _, row in df_g.iterrows():
        ger_ofertas_detalhe.append({
            'polo': str(row.get(c_polo, '')), 'categoria': str(row.get(c_cat, '')),
            'ordem': str(row.get('_ORDEM_G', '')), 'pratica': str(row.get('_PRATICA_G', '')),
            'curso': str(row.get(c_curso, '')),
            'tutor': str(row.get(c_tutor, '')) if pd.notna(row.get(c_tutor)) else '',
            'gerenciado': bool(row.get('_GERENCIADO', False)),
            'tem_agenda': bool(row.get('_TEM_AGENDA', False)),
            'alunos_mat': int(row.get('_ALUNOS_MAT', 0)), 'alunos_agend': int(row.get('_QTD_ALUN', 0)),
            'capacidade': int(row.get('_CAPA', 0)),
        })
    print(f"[{ts()}] Gerenciamento: {len(ger_polo)} polos, {len(ger_cat)} cats, {len(ger_ordem)} ordens")
    return {
        'ger_kpis': ger_kpis, 'ger_polo': ger_polo, 'ger_cat': ger_cat,
        'ger_ordem': ger_ordem, 'ger_contratacao': ger_contratacao,
        'ger_agendas': ger_agendas, 'ger_ofertas': ger_ofertas_detalhe,
    }





def carregar_alunos_hub(path_csv):
    """
    Lê Relatorio_alunos_por_hub.csv e retorna dict com matrículas distintas
    por polo e por categoria — substitui a contagem inflacionada do GIOCONDA.
    """
    import unicodedata as _ud, re as _re
    if not path_csv or not os.path.isfile(path_csv):
        print(f"[{ts()}] Alunos hub: arquivo não encontrado ({path_csv})")
        return None
    print(f"[{ts()}] Lendo alunos por hub: {os.path.basename(path_csv)}")
    for enc in ['latin-1', 'utf-8', 'cp1252']:
        try:
            df = pd.read_csv(path_csv, sep=';', encoding=enc, dtype=str)
            if 'MATRICULA' in df.columns: break
        except: continue
    else:
        print(f"[{ts()}] ERRO: não foi possível ler {path_csv}")
        return None

    # Apenas matrículas confirmadas
    if 'SITUACAO_SEMESTRE' in df.columns:
        df = df[df['SITUACAO_SEMESTRE'].str.strip() == 'Matrícula Confirmada'].copy()

    def _norm(s):
        s = _ud.normalize('NFD', str(s or '').upper().strip())
        s = ''.join(c for c in s if _ud.category(c) != 'Mn')
        s = _re.sub(r'^LAP\s*[-–]\s*', '', s).strip()
        return _re.sub(r'\s+', ' ', s)

    # Mapear GRUPO_HUB → nossas categorias
    GRUPO_CAT = {
        'MULTIDISCIPLINAR II':        'ENF-INS (Multidisciplinar II)',
        'MULTIDISCIPLINAR I':         'BIO-FAR (Multidisciplinar I)',
        'MULTIDISCIPLINAR III':       'BIO-FISIO-EST-TO (Multidisciplinar III)',
        'ENGMAKER+QUIMICA E FISICA':  'QUÍMICA E FÍSICA',
        'ENGMAKER':                   'ENGMAKER',
        'MULTIDISCIPLINAR IV':        'NUTRI (Multidisciplinar IV)',
    }
    def _grupo_para_cat(g):
        gn = _norm(g)
        for k, v in GRUPO_CAT.items():
            if _norm(k) in gn or gn in _norm(k): return v
        return g

    df['_POLO_NORM'] = df['POLO_HUB'].apply(_norm)
    df['_CAT']       = df['GRUPO_HUB'].apply(_grupo_para_cat)

    total_distintos = df['MATRICULA'].nunique()
    print(f"[{ts()}] Matrículas DISTINTAS (ativos): {total_distintos:,}")

    # Por polo (chave normalizada)
    por_polo = (df.groupby('_POLO_NORM')['MATRICULA']
                  .nunique().to_dict())

    # Por polo × categoria
    por_polo_cat = {}
    for (polo, cat), grp in df.groupby(['_POLO_NORM', '_CAT']):
        por_polo_cat[f"{polo}||{cat}"] = int(grp['MATRICULA'].nunique())

    # Por categoria (totais)
    por_cat = (df.groupby('_CAT')['MATRICULA']
                 .nunique().to_dict())

    # ── Mapear TUTOR_PRATICA → subcurso para Multi 3 ────────────────────
    tutor_subcurso = {}  # nome_norm → 'Fisio'/'T.Oc'/'Est'
    if 'TUTOR_PRATICA' in df.columns and 'DISCIPLINA' in df.columns and 'GRUPO_HUB' in df.columns:
        import re as _re
        from collections import Counter as _Counter
        _FISIO = ['FISIOTERAPIA','CINESIOTERAPIA','ELETROTERM','CARDIORRESPIR',
                  'PROTESE','ORTESE','RECURSOS TERAPEUTICOS','MOVIMENTO FUNCIONAL',
                  'AVALIACAO FISICO','REABILITACAO','NEUROFUNC','ORTOPEDIC','RESPIRATORIA']
        _TO    = ['TERAPIA OCUPACIONAL','PSICOMOTRICIDADE','INTEGRACAO SENSORIAL',
                  'TRANSTORNOS MENTAIS','COMPORTAMENTO HUMANO','VIDA DIARIA','TRABALHO EM GRUPO']
        _EST   = ['ESTETICA','COSMETOLOGIA','BIOMEDICINA ESTETICA','PIGMENTAC',
                  'DEPILAC','FACIAL CORPORAL','MICROAGULH']
        def _classif_disc(d):
            d2 = _norm(d) if d else ''
            if any(k in d2 for k in _FISIO): return 'Fisio'
            if any(k in d2 for k in _TO):    return 'T.Oc'
            if any(k in d2 for k in _EST):   return 'Est'
            return None
        def _norm_tutor(s):
            s = _re.sub(r'\s*\(\d+\)\s*$', '', str(s or '')).strip()
            return _norm(s)
        df3 = df[df['GRUPO_HUB'].str.upper().str.contains('MULTIDISCIPLINAR III|MULTI.*3|BIO-FISIO', na=False)].copy()
        df3 = df3[df3['TUTOR_PRATICA'].notna() & (df3['TUTOR_PRATICA'].astype(str).str.strip().str.upper() != 'NAN')]
        df3['_sub'] = df3['DISCIPLINA'].apply(_classif_disc)
        df3['_tnorm'] = df3['TUTOR_PRATICA'].apply(_norm_tutor)
        for tutor, grp in df3[df3['_sub'].notna()].groupby('_tnorm'):
            subs = list(grp['_sub'])
            if subs:
                tutor_subcurso[tutor] = _Counter(subs).most_common(1)[0][0]
        print(f"[{ts()}] Subcursos Multi 3 mapeados: {len(tutor_subcurso)} tutores")

    return {
        'total_distintos': int(total_distintos),
        'por_polo': {k: int(v) for k, v in por_polo.items()},
        'por_polo_cat': por_polo_cat,
        'por_cat': {k: int(v) for k, v in por_cat.items()},
        'tutor_subcurso': tutor_subcurso,  # Multi 3: nome_tutor → Fisio/T.Oc/Est
    }

def gerar_html(dados):
    saida = os.path.join(SCRIPT_DIR, "saida")
    os.makedirs(saida, exist_ok=True)
    output = os.path.join(saida, "dashboard.html")
    tmpl   = os.path.join(SCRIPT_DIR, "template_dashboard.html")
    with open(tmpl, encoding='utf-8') as f: html = f.read()
    html = html.replace("'DATA_GOES_HERE'", json.dumps(dados, ensure_ascii=False))
    html = html.replace("TIMESTAMP_GOES_HERE", dados['gerado_em'])
    with open(output, 'w', encoding='utf-8') as f: f.write(html)
    print(f"[{ts()}] Salvo: {output}")
    return output


def modo_watch(p1, p2):
    print(f"[{ts()}] Monitorando a cada 30s — feche a janela para parar")
    mods = {p1: 0.0, p2: 0.0}
    def loop():
        while True:
            try:
                mudou = any(os.path.getmtime(a) != mods[a] for a in mods if os.path.isfile(a))
                if mudou:
                    for a in mods:
                        if os.path.isfile(a): mods[a] = os.path.getmtime(a)
                    print(f"[{ts()}] Mudança detectada, atualizando...")
                    gerar_html(processar(p1, p2))
            except Exception as e: print(f"[{ts()}] Erro: {e}")
            time.sleep(30)
    threading.Thread(target=loop, daemon=True).start()
    try:
        while True: time.sleep(1)
    except KeyboardInterrupt: print(f"\n[{ts()}] Encerrado.")


if __name__ == '__main__':
    print()
    print(" Verificando arquivos...")
    print()
    p1, p2, tmpl, p3, p4, p5 = verificar_e_localizar()
    if not p1 or not p2 or not os.path.isfile(tmpl):
        print()
        print(" Coloque as planilhas na pasta planilhas\\")
        print(" e tente novamente.")
        print()
        if '--sem-browser' not in sys.argv:
            input(" Pressione Enter para sair...")
        sys.exit(1)
    print()
    dados = processar(p1, p2)
    if p4:
        try:
            lotacao = carregar_lotacao(p4)
            dados = enriquecer_tutores(dados, lotacao)
        except Exception as e:
            print(f"[{ts()}] AVISO: Erro ao processar lotação: {e}")
            dados['alunos_por_curso'] = []
    else:
        dados['alunos_por_curso'] = []
    # PATCH 2: tem_lotacao baseado em dados reais (CH > 0 em pelo menos 1 tutor)
    _ch_ok = sum(1 for t in dados.get('tutores', []) if t.get('ch_semanal') and t['ch_semanal'] > 0)
    dados['tem_lotacao'] = _ch_ok > 0
    print(f"[{ts()}] tem_lotacao={dados['tem_lotacao']} ({_ch_ok} tutores com CH SEMANAL)")
    if p3:
        try:
            ger_dados = processar_gerenciamento(p3)
            dados.update(ger_dados)
            dados['tem_gerenciamento'] = True
            # Enriquecer ger_ofertas com ch_semanal (join por nome normalizado)
            def _norm_nome(s):
                import unicodedata
                s = str(s or '').lower().split('(')[0].strip()
                s = unicodedata.normalize('NFD', s)
                s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
                return ' '.join(s.split())
            def _nome_fl(s):
                pts = _norm_nome(s).split()
                return (pts[0] + ' ' + pts[-1]) if len(pts) >= 2 else _norm_nome(s)
            # Mapear ch_semanal por nome completo E por primeiro+último nome
            _ch_map = {}; _ch_map_fl = {}
            for t in dados.get('tutores', []):
                if t.get('ch_semanal') and t.get('n'):
                    _ch_map[_norm_nome(t['n'])] = t['ch_semanal']
                    _ch_map_fl[_nome_fl(t['n'])] = t['ch_semanal']
            # Injetar ch_semanal em cada oferta
            enr = 0
            # Pré-computar lista de (nome_normalizado, nome_fl, ch) para lookup rápido
            _lot_list = [(k, _nome_fl(k), v) for k, v in _ch_map.items()]

            for oferta in dados.get('ger_ofertas', []):
                tutor = oferta.get('tutor', '')
                if not tutor or oferta.get('ch_semanal'): continue
                tn = _norm_nome(tutor); tfl = _nome_fl(tutor)

                # Match 1: exato ou FL
                ch = _ch_map.get(tn) or _ch_map.get(tfl) or _ch_map_fl.get(tn) or _ch_map_fl.get(tfl)

                # Match 2: tokens do GIOCONDA presentes no nome da lotação
                if not ch:
                    _tokens = [t for t in tfl.split() if len(t) > 2]
                    if len(_tokens) >= 2:
                        for lot_n, lot_fl, lot_ch in _lot_list:
                            if all(tok in lot_n for tok in _tokens) or all(tok in lot_fl for tok in _tokens):
                                ch = lot_ch; break

                # Match 3: tokens da LOTAÇÃO presentes no nome do GIOCONDA (inverso)
                if not ch:
                    for lot_n, lot_fl, lot_ch in _lot_list:
                        lot_tokens = [t for t in lot_fl.split() if len(t) > 2]
                        if len(lot_tokens) >= 2 and all(tok in tn for tok in lot_tokens):
                            ch = lot_ch; break

                if ch:
                    oferta['ch_semanal'] = ch; enr += 1
            print(f"[{ts()}] CH enriquecida: {enr}/{len(dados.get('ger_ofertas',[]))} ofertas")
        except Exception as e:
            print(f"[{ts()}] AVISO: Erro ao processar gerenciamento: {e}")
            import traceback; traceback.print_exc()
            dados['tem_gerenciamento'] = False
    else:
        dados['tem_gerenciamento'] = False
    # ── ALUNOS HUB: matrículas distintas ──────────────────────────────────────
    if p5:
        try:
            alunos_hub = carregar_alunos_hub(p5)
            if alunos_hub:
                dados['alunos_hub'] = alunos_hub
                # Sobrescrever total_alunos_matriculados nos ger_kpis
                if 'ger_kpis' in dados:
                    dados['ger_kpis']['total_alunos_matriculados'] = alunos_hub['total_distintos']
                    dados['ger_kpis']['alunos_mat_fonte'] = 'hub_csv'
                    print(f"[{ts()}] KPI alunos substituído: {alunos_hub['total_distintos']:,} (matrículas distintas)")
        except Exception as e:
            print(f"[{ts()}] AVISO: erro ao ler alunos hub: {e}")
    else:
        print(f"[{ts()}] INFO: Relatorio_alunos_por_hub.csv não encontrado — usando contagem GIOCONDA")
    html = gerar_html(dados)
    if '--sem-browser' not in sys.argv:
        print(f"[{ts()}] Abrindo navegador...")
        webbrowser.open(Path(html).as_uri())
    if WATCH_MODE: modo_watch(p1, p2)
    else: print(f"[{ts()}] Concluído!")
