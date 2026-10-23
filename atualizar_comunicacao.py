"""Gera tutores_comunicacao.json a partir do Engajômetro (aba Acompanhamento).

Uso:  python atualizar_comunicacao.py "C:/Users/leona/Downloads/Engajometro.xlsx"

O JSON alimenta o bloco "Acompanhamento de comunicação" da ficha do tutor
(PATCH 13 em processar.py). Formato de cada item: nome, polo, atencao, resumo,
tentativas_contato (lista), ultimo_contato, observacao + categoria, desligamento,
multiplicador e data_registro.

Cuidados com o dado da planilha (preenchida à mão por vários multiplicadores):
- Datas com dia <= 12 que o Excel guardou como data foram lidas no formato
  americano (ex.: "01/07/2026" virou 7 de janeiro). Para essas, o texto
  digitado é recuperado como MM/DD -> DD/MM. Com dia > 12 não há troca possível.
- "Tentativas" pode ser lista de datas, um número ou texto livre.
- Linhas do modelo ("Selecione o curso", sem nome) são ignoradas.
- Mesma pessoa em mais de uma linha (nome curto e completo, mesmo polo): fica
  o registro mais recente, com o nome completo.
"""
import datetime as dt
import json
import os
import re
import sys
import unicodedata

from openpyxl import load_workbook

SAIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'tutores_comunicacao.json')


def _txt(v):
    if v is None:
        return ''
    if isinstance(v, (dt.datetime, dt.date)):
        # Excel leu DD/MM como MM/DD quando o dia era <= 12 (ver docstring)
        return f'{v.month:02d}/{v.day:02d}/{v.year}' if v.day <= 12 else f'{v.day:02d}/{v.month:02d}/{v.year}'
    return re.sub(r'\s+', ' ', str(v)).strip()


def _norm(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode('ascii')
    return ' '.join(s.upper().split())


def _data_ord(s):
    m = re.match(r'(\d{2})/(\d{2})/(\d{4})', s or '')
    return f'{m.group(3)}{m.group(2)}{m.group(1)}' if m else ''


def _tentativas(v):
    t = _txt(v)
    if not t:
        return []
    if re.fullmatch(r'\d+', t):
        return [f'{t} tentativa' + ('s' if t != '1' else '')]
    if re.fullmatch(r'\d{1,2}/\d{1,2}(/\d{2,4})?(\s*,\s*\d{1,2}/\d{1,2}(/\d{2,4})?)*', t):
        return [x.strip() for x in t.split(',')]
    return [t]


def main(caminho):
    ws = load_workbook(caminho, data_only=True)['Acompanhamento']
    cab = [_norm(_txt(c.value)) for c in ws[1]]
    col = lambda nome: cab.index(nome)
    itens = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        nome = _txt(r[col('NOME DO TUTOR')])
        if not nome:
            continue
        deslig = r[col('DESLIGAMENTO')]
        itens.append({
            'nome': nome,
            'polo': _txt(r[col('POLO')]),
            'atencao': _norm(_txt(r[col('REQUER ATENCAO URGENTE')])) == 'SIM',
            'resumo': _txt(r[col('RESUMO DA SITUACAO')]) or None,
            'tentativas_contato': _tentativas(r[col('TENTATIVAS DE CONTATO')]),
            'ultimo_contato': _txt(r[col('ULTIMO CONTATO')]) or None,
            'observacao': _txt(r[col('OBSERVACAO')]) or None,
            'categoria': _txt(r[col('CATEGORIA')]) or None,
            'desligamento': {True: 'Sim', False: 'Não'}.get(deslig, _txt(deslig) or None),
            'multiplicador': _txt(r[col('MULTIPLICADOR RESPONSAVEL')]) or None,
            'data_registro': _txt(r[col('DATA DO REGISTRO')]) or None,
        })

    # mesma pessoa em 2 linhas: 1º e último nome iguais + mesma cidade do polo
    def chave(i):
        toks = _norm(i['nome']).split()
        cidade = _norm(re.split(r'[/-]', i['polo'])[0]) if i['polo'] else ''
        return (toks[0], toks[-1], cidade)
    finais = {}
    for i in itens:
        k = chave(i)
        a = finais.get(k)
        if a is None:
            finais[k] = i
            continue
        novo, velho = (i, a) if _data_ord(i['data_registro']) >= _data_ord(a['data_registro']) else (a, i)
        if len(velho['nome']) > len(novo['nome']):
            novo['nome'] = velho['nome']
        print(f'  mesclado: "{velho["nome"]}" + "{novo["nome"]}" ({novo["polo"]}) -> fica o registro de {novo["data_registro"]}')
        finais[k] = novo
    lista = list(finais.values())
    with open(SAIDA, 'w', encoding='utf-8') as f:
        json.dump(lista, f, ensure_ascii=False, indent=1)
        f.write('\n')
    print(f'{len(lista)} tutores gravados em {SAIDA} ({sum(i["atencao"] for i in lista)} com atenção urgente)')


if __name__ == '__main__':
    main(sys.argv[1])
