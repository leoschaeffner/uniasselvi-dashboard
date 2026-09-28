"""Gera planilhas/REMESSAS_INSUMOS_template.xlsx (gitignored) para a equipe
de laboratórios preencher — fonte p10/URL_REMESSAS do processar.py (PATCH 178).
Aba 'Registro' (cabeçalho + 3 linhas de EXEMPLO) e aba oculta 'Listas' que
só alimenta os dropdowns (STATUS, TRANSPORTADORA, CATEGORIA, POLO).
Nenhuma coluna de PII (destinatário/telefone/e-mail/endereço) de propósito.
Uso: python tests/tools/gerar_template_remessas.py"""
import os, sys, json
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, RAIZ)
import processar as P  # listas canônicas (status/categorias) vêm do próprio pipeline

SAIDA = os.path.join(RAIZ, 'planilhas', 'REMESSAS_INSUMOS_template.xlsx')
COLUNAS = ['Polo', 'Categoria de Laboratório', 'Item/Kit', 'Quantidade', 'Unidade', 'Data de Envio',
           'Previsão de Entrega', 'Data de Entrega', 'Transportadora', 'Código de Rastreio', 'Status', 'Observação']
LARGURAS = [34, 26, 30, 12, 10, 15, 18, 15, 18, 20, 14, 40]

with open(os.path.join(RAIZ, 'contatos_por_polo.json'), encoding='utf-8') as f:
    POLOS = sorted(json.load(f).keys())  # só as chaves (nomes de polo), nunca os contatos
TRANSP = [t['nome'] for t in P._carregar_transportadoras()] or ['Correios']

wb = Workbook()
ws = wb.active
ws.title = 'Registro'
ws.append(COLUNAS)
fill = PatternFill('solid', fgColor='0F766E')
borda = Border(*(Side(style='thin', color='CBD5E1'),) * 4)
for i, c in enumerate(ws[1], 1):
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = fill
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = LARGURAS[i - 1]
ws.row_dimensions[1].height = 30
ws.freeze_panes = 'A2'
h = date.today()
exemplos = [
    [POLOS[0], 'Enfermagem', 'Kit de insumos', 2, 'kit', h - timedelta(days=5), h + timedelta(days=3), None,
     'Correios', 'AA123456789BR', 'Em trânsito', 'EXEMPLO — apagar'],
    [POLOS[1], 'EngeMaker', 'Kit de EPIs', 1, 'kit', h - timedelta(days=12), h - timedelta(days=4), h - timedelta(days=6),
     'Jadlog', 'JD0000000000', 'Entregue', 'EXEMPLO — apagar'],
    [POLOS[2], 'Multidisciplinar I', 'Reposição de reagentes', 3, 'caixa', None, h + timedelta(days=8), None,
     '', '', 'Preparando', 'EXEMPLO — apagar'],
]
for r in exemplos:
    ws.append(r)
for row in ws.iter_rows(min_row=2, max_row=200):
    for c in row:
        c.border = borda
    for idx in (5, 6, 7):  # F, G, H = datas
        row[idx].number_format = 'DD/MM/YYYY'
for row in ws.iter_rows(min_row=2, max_row=4):
    for c in row:
        c.font = Font(italic=True, color='94A3B8')

# aba oculta Listas (dropdowns)
ls = wb.create_sheet('Listas')
listas = [('STATUS', P._REM_STATUS), ('TRANSPORTADORA', TRANSP), ('CATEGORIA', P._VIST_CATEGORIAS), ('POLO', POLOS)]
for j, (nome, itens) in enumerate(listas, 1):
    ls.cell(row=1, column=j, value=nome).font = Font(bold=True)
    for i, v in enumerate(itens, 2):
        ls.cell(row=i, column=j, value=v)
ls.sheet_state = 'hidden'

def dropdown(col_letra, lista_col, n, estrito):
    ref = f"=Listas!${get_column_letter(lista_col)}$2:${get_column_letter(lista_col)}${n + 1}"
    dv = DataValidation(type='list', formula1=ref, allow_blank=True, showErrorMessage=estrito)
    ws.add_data_validation(dv)
    dv.add(f'{col_letra}2:{col_letra}500')

dropdown('A', 4, len(POLOS), True)              # Polo
dropdown('B', 3, len(P._VIST_CATEGORIAS), True)  # Categoria
dropdown('I', 2, len(TRANSP), False)             # Transportadora (aceita outra; sem link)
dropdown('K', 1, len(P._REM_STATUS), True)       # Status
dv_dt = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True,
                       showErrorMessage=True, errorTitle='Data inválida', error='Use uma data (DD/MM/AAAA).')
ws.add_data_validation(dv_dt)
for col in 'FGH':
    dv_dt.add(f'{col}2:{col}500')
dv_q = DataValidation(type='whole', operator='greaterThan', formula1='0', allow_blank=True,
                      showErrorMessage=True, error='Quantidade deve ser um número inteiro maior que zero.')
ws.add_data_validation(dv_q)
dv_q.add('D2:D500')

os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
wb.save(SAIDA)
print(f"Template salvo: {SAIDA}")
print(f"  aba Registro: {len(COLUNAS)} colunas, {len(exemplos)} linhas de exemplo; aba oculta Listas: "
      + ', '.join(f'{n}={len(v)}' for n, v in listas))
